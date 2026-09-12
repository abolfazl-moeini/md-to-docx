"""Comprehensive tests for finilize.v3.md requirements (V3-01 to V3-08)."""

import shutil
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from xml.etree import ElementTree as ET
import pytest
from click.testing import CliRunner
from docx import Document

from md_to_docx import Template, convert_markdown_to_docx
from md_to_docx.cli import main
from md_to_docx.fonts_embed import obfuscate_font_data, validate_font_embedding
from md_to_docx.mermaid import ConvertError
from md_to_docx.options import (
    DEFAULT_FONT_FAMILY,
    DEFAULT_LATIN_FONT,
    DEFAULT_CODE_FONT,
    GeneratorOptions,
    resolve_effective_direction,
)
from md_to_docx.renderer import DocxRenderer
from scripts.package_oracle import run_package_oracle


# ============================================================================
# V3-01: Direction Precedence and RTL/LTR Formatting
# ============================================================================

def test_v3_01_direction_precedence_hierarchy():
    """Verify strict precedence: option > meta dir > meta lang > template."""
    # 1. Option wins over everything
    assert resolve_effective_direction(option_direction="ltr", meta_direction="rtl", template_direction="rtl", meta_lang="fa") == "ltr"
    assert resolve_effective_direction(option_direction="rtl", meta_direction="ltr", template_direction="ltr", meta_lang="en") == "rtl"

    # 2. Meta direction wins over meta lang and template when option is auto/None
    assert resolve_effective_direction(option_direction="auto", meta_direction="ltr", template_direction="rtl", meta_lang="fa") == "ltr"
    assert resolve_effective_direction(option_direction=None, meta_direction="rtl", template_direction="ltr", meta_lang="en") == "rtl"

    # 3. Meta lang wins over template when option is auto/None and meta dir is None
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="rtl", meta_lang="en") == "ltr"
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="rtl", meta_lang="en-US") == "ltr"
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="ltr", meta_lang="fa") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="ltr", meta_lang="fa-IR") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="ltr", meta_lang="ar") == "rtl"

    # 4. Template direction is fallback
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="ltr", meta_lang=None) == "ltr"
    assert resolve_effective_direction(option_direction="auto", meta_direction=None, template_direction="rtl", meta_lang=None) == "rtl"
    assert resolve_effective_direction(option_direction=None, meta_direction=None, template_direction="rtl", meta_lang=None) == "rtl"


def test_v3_01_red_test_dir_ltr_on_rtl_template(tmp_path):
    """dir: ltr in metadata on purple_book (RTL template) must produce LTR document and styling."""
    md = """---
dir: ltr
lang: en-US
---

# Introduction to Algorithms

This is an English paragraph describing sorting algorithms.

- First item
- Second item
"""
    out = tmp_path / "ltr_on_rtl.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
        settings_xml = z.read("word/settings.xml").decode("utf-8") if "word/settings.xml" in z.namelist() else ""

    # Document bidi must be absent or 0 in settings/document
    assert '<w:bidi w:val="1"/>' not in settings_xml
    assert '<w:bidi/>' not in settings_xml

    # Headings and body paragraphs should not have w:bidi
    doc_root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for p in doc_root.findall(f".//{W}p"):
        texts = [t.text for t in p.iter(f"{W}t") if t.text]
        full_text = "".join(texts)
        if "Introduction to Algorithms" in full_text or "This is an English" in full_text:
            bidi_el = p.find(f".//{W}bidi")
            assert bidi_el is None or bidi_el.get(f"{W}val") == "0"


def test_v3_01_numbers_urls_sql_in_rtl_text(tmp_path):
    """In RTL text, numbers, URLs, and inline code must have w:rtl=0 and w:cs=0."""
    md = """---
lang: fa-IR
dir: rtl
---

نسخه جدید در سال 2024 با شماره 1.4.1 منتشر شد.

برای دانلود به آدرس https://example.com/download مراجعه کنید.

دستور `SELECT * FROM users WHERE id = 1` را اجرا کنید.
"""
    out = tmp_path / "inline_ltr_in_rtl.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    # Verify runs containing 2024 or 1.4.1 have w:rtl="0" or no w:rtl and w:cs="0"
    for r in root.findall(f".//{W}r"):
        t_el = r.find(f"{W}t")
        if t_el is not None and t_el.text:
            txt = t_el.text
            if "2024" in txt or "1.4.1" in txt:
                rtl_el = r.find(f".//{W}rtl")
                assert rtl_el is None or rtl_el.get(f"{W}val") == "0"
            elif "https://example.com/download" in txt:
                rtl_el = r.find(f".//{W}rtl")
                assert rtl_el is None or rtl_el.get(f"{W}val") == "0"
            elif "SELECT" in txt or "users" in txt:
                rtl_el = r.find(f".//{W}rtl")
                assert rtl_el is None or rtl_el.get(f"{W}val") == "0"


def test_v3_01_mixed_heading_runs(tmp_path):
    """Mixed Persian and Latin heading runs must preserve individual script tags."""
    md = """---
lang: fa-IR
dir: rtl
---

## فصل دوم: معماری Microservices در کلاستر Kubernetes
"""
    out = tmp_path / "mixed_heading.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    found_microservices = False
    found_persian = False
    for r in root.findall(f".//{W}r"):
        t = r.find(f"{W}t")
        if t is not None and t.text:
            if "Microservices" in t.text or "Kubernetes" in t.text:
                found_microservices = True
                rtl_el = r.find(f".//{W}rtl")
                assert rtl_el is None or rtl_el.get(f"{W}val") == "0"
            if "معماری" in t.text or "کلاستر" in t.text:
                found_persian = True
                rtl_el = r.find(f".//{W}rtl")
                assert rtl_el is not None and rtl_el.get(f"{W}val") in ("1", None)

    assert found_microservices and found_persian


def test_v3_01_footnote_rtl_and_ltr(tmp_path):
    """Footnotes in RTL and LTR contexts must set proper paragraph bidi and alignment."""
    md = """---
lang: fa-IR
dir: rtl
---

متن دارای پاورقی فارسی[^fn1] و انگلیسی[^fn2].

[^fn1]: این یک توضیح فارسی در پاورقی است.
[^fn2]: This is an English footnote explanation.
"""
    out = tmp_path / "footnotes_bidi.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        assert "word/footnotes.xml" in z.namelist()
        fn_xml = z.read("word/footnotes.xml").decode("utf-8")

    root = ET.fromstring(fn_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    found_fa_fn = False
    found_en_fn = False
    for p in root.findall(f".//{W}p"):
        texts = [t.text for t in p.iter(f"{W}t") if t.text]
        full_text = "".join(texts)
        if "این یک توضیح فارسی" in full_text:
            found_fa_fn = True
            bidi_el = p.find(f".//{W}bidi")
            assert bidi_el is not None and bidi_el.get(f"{W}val") in ("1", None)
        elif "This is an English footnote" in full_text:
            found_en_fn = True
            bidi_el = p.find(f".//{W}bidi")
            # Pure English footnote in LTR/pure latin resolves bidi=0
            assert bidi_el is None or bidi_el.get(f"{W}val") == "0"

    assert found_fa_fn and found_en_fn


# ============================================================================
# V3-02: Font Options and Typography
# ============================================================================

def test_v3_02_font_options_probe(tmp_path):
    """GeneratorOptions font probe verifies heading_font, font_family, latin_font, code_font."""
    md = """---
lang: fa-IR
dir: rtl
---

# تیتر با فونت ویژه

متن بدنه به فارسی و English sample.

```python
x = 42
```
"""
    opts = GeneratorOptions(
        font_family="ProbeBodyFont",
        heading_font="ProbeHeadingFont",
        latin_font="ProbeLatinFont",
        code_font="ProbeCodeFont",
    )
    out = tmp_path / "probe_fonts.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", options=opts, overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    assert "ProbeBodyFont" in doc_xml
    assert "ProbeHeadingFont" in doc_xml
    assert "ProbeLatinFont" in doc_xml
    assert "ProbeCodeFont" in doc_xml


def test_v3_02_rfonts_east_asia_and_cs_attributes(tmp_path):
    """Run fonts must properly set ascii, hAnsi, cs, and eastAsia."""
    md = "متن فارسی نمونه.\n"
    out = tmp_path / "rfonts_check.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    found_rfonts = False
    for rFonts in root.findall(f".//{W}rFonts"):
        cs = rFonts.get(f"{W}cs")
        ea = rFonts.get(f"{W}eastAsia")
        if cs == "Vazirmatn":
            found_rfonts = True
            assert ea == "Vazirmatn"

    assert found_rfonts


def test_v3_02_font_family_none_uses_template_without_w05_warning():
    """Explicit options with font_family=None is normal template usage and must NOT emit W05 warning (finilize.v3.md line 114)."""
    doc = Document()
    tmpl = Template.load("purple_book")
    opts = GeneratorOptions(font_family=None)
    renderer = DocxRenderer(doc, tmpl, options=opts)
    assert not any("W05" in w for w in renderer.warnings)
    assert renderer.body_font == tmpl.fonts.get("body", "Vazirmatn")


# ============================================================================
# V3-03: Font Embedding (ECMA-376) and Validation
# ============================================================================

def test_v3_03_font_obfuscation_roundtrip():
    """Verify ECMA-376 TrueType obfuscation XOR key reversible behavior."""
    guid = "{12345678-ABCD-EF01-2345-6789ABCDEF01}"
    original = b"This is a 64-byte sample font payload that tests XOR obfuscation." * 2
    obfuscated = obfuscate_font_data(original, guid)
    # Applying XOR with the same key again restores original bytes
    restored = obfuscate_font_data(obfuscated, guid)
    assert restored == original
    # Obfuscated must differ from original in the first 32 bytes
    assert obfuscated[:32] != original[:32]
    assert obfuscated[32:] == original[32:]


def test_v3_03_font_embedding_success_and_package_oracle(tmp_path):
    """When embed_fonts=True, DOCX contains obfuscated fonts and passes package oracle."""
    md = """---
lang: fa-IR
dir: rtl
---

# آزمون جاسازی فونت

این سند با فونت توکار Vazirmatn ایجاد شده است.
"""
    out = tmp_path / "embedded.docx"
    convert_markdown_to_docx(
        content=md,
        output_path=out,
        template="purple_book",
        embed_fonts=True,
        overwrite=True,
    )

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "word/fonts/font1.odttf" in names
        assert "word/fonts/font2.odttf" in names
        assert "word/fontTable.xml" in names
        assert "word/_rels/fontTable.xml.rels" in names

        # Verify [Content_Types].xml includes odttf and fontTable
        ct_xml = z.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="odttf"' in ct_xml
        assert 'PartName="/word/fontTable.xml"' in ct_xml

        # Verify relationships
        doc_rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        assert 'Target="fontTable.xml"' in doc_rels
        assert 'fontTable' in doc_rels

        ft_rels = z.read("word/_rels/fontTable.xml.rels").decode("utf-8")
        assert 'Target="fonts/font1.odttf"' in ft_rels
        assert 'Target="fonts/font2.odttf"' in ft_rels

        # Verify fontTable.xml has w:font with embedRegular and embedBold
        ft_xml = z.read("word/fontTable.xml").decode("utf-8")
        assert '<w:font w:name="Vazirmatn">' in ft_xml
        assert 'w:embedRegular' in ft_xml
        assert 'w:embedBold' in ft_xml

    # Verify package passes package oracle without errors
    ok, issues = run_package_oracle(out)
    assert ok, f"Package oracle reported issues for embedded DOCX: {issues}"


def test_v3_03_negative_missing_font_file_rejected(tmp_path):
    """embed_fonts=True with a non-existent font file must raise ConvertError."""
    import yaml
    custom_dir = tmp_path / "missing_font_tmpl"
    pb_dir = Path("templates/purple_book")
    shutil.copytree(pb_dir, custom_dir)

    # Clear font_files in config.yaml so Template.load passes
    cfg_path = custom_dir / "config.yaml"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    cfg["font_files"] = {}
    cfg_path.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")

    # Remove font files
    shutil.rmtree(custom_dir / "fonts")

    tmpl = Template.load(custom_dir)
    out = tmp_path / "fail_embed.docx"

    with pytest.raises(ConvertError, match="[Ff]ont.*embedding|[Nn]ot found"):
        convert_markdown_to_docx(
            content="# Hello",
            output_path=out,
            template=tmpl,
            embed_fonts=True,
            overwrite=True,
        )


def test_v3_03_negative_missing_license_rejected(tmp_path):
    """embed_fonts=True with missing OFL.txt license file must raise ConvertError."""
    custom_dir = tmp_path / "missing_license_tmpl"
    pb_dir = Path("templates/purple_book")
    shutil.copytree(pb_dir, custom_dir)

    # Remove license file
    lic_file = custom_dir / "fonts" / "OFL.txt"
    if lic_file.exists():
        lic_file.unlink()

    tmpl = Template.load(custom_dir)
    out = tmp_path / "fail_lic.docx"

    with pytest.raises(ConvertError, match="[Ll]icense|[Ee]mbedding"):
        convert_markdown_to_docx(
            content="# Hello",
            output_path=out,
            template=tmpl,
            embed_fonts=True,
            overwrite=True,
        )


# ============================================================================
# CLI Option Verification
# ============================================================================

def test_v3_cli_options_file_conversion(tmp_path):
    """CLI supports --font, --embed-fonts, --direction, --text-align."""
    runner = CliRunner()
    in_file = tmp_path / "cli_input.md"
    in_file.write_text("# آزمایش CLI\n\nمتن بدنه.\n", encoding="utf-8")
    out_file = tmp_path / "cli_output.docx"

    result = runner.invoke(
        main,
        [
            "convert",
            str(in_file),
            "-o",
            str(out_file),
            "--template",
            "purple_book",
            "--font",
            "Vazirmatn",
            "--embed-fonts",
            "--direction",
            "rtl",
            "--text-align",
            "start",
            "-f",
        ],
    )
    assert result.exit_code == 0, f"CLI exited with {result.exit_code}: {result.output}"
    assert "font: Vazirmatn" in result.output
    assert "embedded: yes" in result.output
    assert "direction: rtl" in result.output


def test_v3_cli_options_stdin_conversion(tmp_path):
    """CLI supports reading from stdin with new GeneratorOptions flags."""
    runner = CliRunner()
    out_file = tmp_path / "stdin_output.docx"

    result = runner.invoke(
        main,
        [
            "convert",
            "-",
            "-o",
            str(out_file),
            "--template",
            "purple_book",
            "--font",
            "Vazirmatn",
            "--direction",
            "rtl",
            "-f",
        ],
        input="# عنوان از ورودی استاندارد\n\nمتن فارسی.\n",
    )
    assert result.exit_code == 0, f"CLI stdin exited with {result.exit_code}: {result.output}"
    assert out_file.exists()


# ============================================================================
# Additional Comprehensive V3 Tests
# ============================================================================

def test_v3_01_narrative_direction_heuristic(tmp_path):
    """Persian narrative text without metadata on LTR template resolves to RTL (finilize.v3.md Section 1.1)."""
    md = """# مبانی یادگیری ماشین

این یک متن مقدماتی درباره الگوریتم‌های هوش مصنوعی و بهینه‌سازی داده‌ها است.

- رگرسیون خطی
- شبکه‌های عصبی
"""
    # Create an LTR custom template
    import yaml
    custom_dir = tmp_path / "ltr_tmpl"
    shutil.copytree("templates/purple_book", custom_dir)
    cfg_path = custom_dir / "config.yaml"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    cfg["direction"] = "ltr"
    cfg_path.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")

    tmpl = Template.load(custom_dir)
    out = tmp_path / "narrative_rtl.docx"
    convert_markdown_to_docx(content=md, output_path=out, template=tmpl, overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    # Heading and body paragraph should have w:bidi set
    found_heading = False
    for p in root.findall(f".//{W}p"):
        texts = [t.text for t in p.iter(f"{W}t") if t.text]
        full_text = "".join(texts)
        if "مبانی یادگیری ماشین" in full_text:
            found_heading = True
            bidi_el = p.find(f".//{W}bidi")
            assert bidi_el is not None and bidi_el.get(f"{W}val") in ("1", None)
    assert found_heading


def test_v3_01_code_comments_do_not_flip_narrative_direction(tmp_path):
    """Persian comments inside code blocks do not make an English document RTL (finilize.v3.md Section 1.1)."""
    md = """# Binary Search Implementation

This document describes the binary search algorithm implemented in Python.

```python
# این یک کامنت فارسی طولانی است که نباید جهت سند را به راست به چپ تغییر دهد
def binary_search(arr, target):
    pass
```
"""
    out = tmp_path / "code_comment_ltr.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)

    with zipfile.ZipFile(out) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")

    root = ET.fromstring(doc_xml)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for p in root.findall(f".//{W}p"):
        texts = [t.text for t in p.iter(f"{W}t") if t.text]
        full_text = "".join(texts)
        if "Binary Search Implementation" in full_text:
            bidi_el = p.find(f".//{W}bidi")
            assert bidi_el is None or bidi_el.get(f"{W}val") == "0"


def test_v3_01_metadata_alias_conflict_warning(tmp_path):
    """Conflicting dir/direction or lang/language metadata emits diagnostic warning."""
    md = """---
dir: ltr
direction: rtl
lang: fa
language: en
---

# Test Conflict
"""
    warnings_list = []
    out = tmp_path / "conflict.docx"
    convert_markdown_to_docx(
        content=md,
        output_path=out,
        template="purple_book",
        warnings=warnings_list,
        overwrite=True,
    )
    assert any("W06" in w and "Conflicting direction metadata" in w for w in warnings_list)
    assert any("W07" in w and "Conflicting language metadata" in w for w in warnings_list)


def test_v3_02_heading_font_precedence():
    """Verify heading font precedence: override > font_family > template heading > body_font."""
    doc = Document()
    tmpl = Template.load("purple_book")

    # 1. heading_font override wins over font_family
    opts1 = GeneratorOptions(font_family="Sahel", heading_font="Shabnam")
    r1 = DocxRenderer(doc, tmpl, options=opts1)
    assert r1.body_font == "Sahel"
    assert r1.heading_font == "Shabnam"

    # 2. font_family alone applies to body AND heading when heading_font is None
    opts2 = GeneratorOptions(font_family="Sahel")
    r2 = DocxRenderer(doc, tmpl, options=opts2)
    assert r2.body_font == "Sahel"
    assert r2.heading_font == "Sahel"

    # 3. Without options, template heading font is used
    opts3 = GeneratorOptions()
    r3 = DocxRenderer(doc, tmpl, options=opts3)
    assert r3.body_font == "Vazirmatn"
    assert r3.heading_font == "Vazirmatn"


def test_v3_02_central_validation():
    """GeneratorOptions central validation rejects invalid values and normalizes aliases."""
    # Unknown key in from_dict
    with pytest.raises(ValueError, match="Unknown generator option"):
        GeneratorOptions.from_dict({"unknown_field": "val"})

    # embed_fonts not boolean
    with pytest.raises(ValueError, match="embed_fonts must be a boolean"):
        GeneratorOptions(embed_fonts="false")  # type: ignore

    # Invalid direction
    with pytest.raises(ValueError, match="Invalid direction 'up'"):
        GeneratorOptions(direction="up")

    # Invalid text_align
    with pytest.raises(ValueError, match="Invalid text_align 'middle'"):
        GeneratorOptions(text_align="middle")

    # Empty font string
    with pytest.raises(ValueError, match="font_family cannot be an empty"):
        GeneratorOptions(font_family="   ")

    # Justify alias normalized to both
    opts = GeneratorOptions(text_align="justify")
    assert opts.text_align == "both"


def test_v3_03_settings_xml_embed_true_type_fonts(tmp_path):
    """When embed_fonts=True, word/settings.xml must have w:embedTrueTypeFonts and saveSubsetFonts stripped."""
    md = "# Embed Settings Test\n\nمتن با فونت توکار.\n"
    out = tmp_path / "settings_check.docx"
    convert_markdown_to_docx(
        content=md,
        output_path=out,
        template="purple_book",
        embed_fonts=True,
        overwrite=True,
    )

    with zipfile.ZipFile(out) as z:
        assert "word/settings.xml" in z.namelist()
        st_xml = z.read("word/settings.xml").decode("utf-8")
        assert "embedTrueTypeFonts" in st_xml
        assert "saveSubsetFonts" not in st_xml


def test_v3_03_fs_type_restricted_rejected(tmp_path):
    """Font with restricted fsType (0x0002) must be rejected by validate_font_embedding."""
    from md_to_docx.fonts_embed import validate_font_embedding
    import yaml

    custom_dir = tmp_path / "restricted_tmpl"
    shutil.copytree("templates/purple_book", custom_dir)

    # Patch fsType in regular font file copy
    reg_font = custom_dir / "fonts" / "Vazirmatn-Regular.ttf"
    data = bytearray(reg_font.read_bytes())
    # Locate OS/2 table and set fsType to 0x0002
    for i in range(len(data) - 4):
        if data[i:i+4] == b"OS/2":
            # Read offset of OS/2 table from directory
            offset = int.from_bytes(data[i+8:i+12], "big")
            # fsType is at offset + 8
            data[offset+8:offset+10] = (2).to_bytes(2, "big")
            break
    reg_font.write_bytes(data)

    tmpl = Template.load(custom_dir)
    valid, err, _, _ = validate_font_embedding(tmpl)
    assert not valid
    assert "restricted licensing" in err.lower()


def test_v3_03_strip_embedded_fonts_when_disabled(tmp_path):
    """When embed_fonts=False, any embedded fonts in intermediate package are stripped."""
    from md_to_docx.fonts_embed import strip_embedded_fonts

    # First build an embedded document
    md = "# Embedded Document\n\nمتن فارسی.\n"
    out = tmp_path / "with_embed.docx"
    convert_markdown_to_docx(
        content=md,
        output_path=out,
        template="purple_book",
        embed_fonts=True,
        overwrite=True,
    )

    with zipfile.ZipFile(out) as z:
        assert "word/fonts/font1.odttf" in z.namelist()

    # Now strip embedded fonts
    strip_embedded_fonts(out)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert not any(n.startswith("word/fonts/") for n in names)
        if "word/settings.xml" in names:
            st = z.read("word/settings.xml").decode("utf-8")
            assert "embedTrueTypeFonts" not in st
        if "word/fontTable.xml" in names:
            ft = z.read("word/fontTable.xml").decode("utf-8")
            assert "embedRegular" not in ft
            assert "embedBold" not in ft


def test_v3_cli_additional_font_flags(tmp_path):
    """CLI supports --heading-font, --latin-font, and --code-font flags."""
    runner = CliRunner()
    in_file = tmp_path / "cli_fonts.md"
    in_file.write_text("# تیتر\n\nمتن لاتین test.\n", encoding="utf-8")
    out_file = tmp_path / "cli_fonts.docx"

    result = runner.invoke(
        main,
        [
            "convert",
            str(in_file),
            "-o",
            str(out_file),
            "--template",
            "purple_book",
            "--font",
            "Sahel",
            "--heading-font",
            "Shabnam",
            "--latin-font",
            "Calibri",
            "--code-font",
            "Consolas",
            "-f",
        ],
    )
    assert result.exit_code == 0, f"CLI exited with {result.exit_code}: {result.output}"
    assert out_file.exists()
    assert "font: Sahel" in result.output


def test_v3_01_bcp47_script_subtags_rtl():
    """BCP-47 language tags with Arabic/Hebrew/Syriac script subtags resolve to RTL."""
    assert resolve_effective_direction(option_direction="auto", meta_lang="az-Arab", template_direction="ltr") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_lang="ku-Arab", template_direction="ltr") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_lang="pa-Arab", template_direction="ltr") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_lang="he-IL", template_direction="ltr") == "rtl"
    assert resolve_effective_direction(option_direction="auto", meta_lang="en-US", template_direction="rtl") == "ltr"


def test_v3_02_heading_font_applied_to_numbered_badge_and_callout():
    """Effective heading_font is applied to heading badges, callout titles, and table headers."""
    from md_to_docx.headings import HeadingInfo

    doc = Document()
    tmpl = Template.load("purple_book")
    opts = GeneratorOptions(font_family="Sahel", heading_font="Shabnam")
    renderer = DocxRenderer(doc, tmpl, options=opts)

    # 1. Numbered heading badge
    tbl = renderer.render_heading(HeadingInfo(level=1, title="عنوان تیتر", number="۱", raw_text="# ۱. عنوان تیتر"))
    cell0 = tbl.cell(0, 0)
    p_badge_xml = cell0.paragraphs[0]._p.xml
    assert 'w:cs="Shabnam"' in p_badge_xml

    # 2. Callout header title
    callout_tbl = renderer.render_callout("note", "عنوان یادداشت", ["متن یادداشت"])
    callout_hdr_xml = callout_tbl.cell(0, 0).paragraphs[0]._p.xml
    assert 'w:cs="Shabnam"' in callout_hdr_xml

    # 3. Table header
    table = renderer.render_table(["ستون اول", "ستون دوم"], [["داده ۱", "داده ۲"]])
    th_xml = table.cell(0, 0).paragraphs[0]._p.xml
    assert 'w:cs="Shabnam"' in th_xml
    td_xml = table.cell(1, 0).paragraphs[0]._p.xml
    assert 'w:cs="Sahel"' in td_xml


def test_v3_03_corrupt_font_rejected_by_validation(tmp_path):
    """Font with corrupt header or missing OS/2 table is rejected by validate_font_embedding."""
    custom_dir = tmp_path / "corrupt_font_tmpl"
    shutil.copytree("templates/purple_book", custom_dir)
    reg_font = custom_dir / "fonts" / "Vazirmatn-Regular.ttf"
    # Overwrite font with invalid 30 zero bytes
    reg_font.write_bytes(b"\x00" * 30)

    tmpl = Template.load(custom_dir)
    valid, err, _, _ = validate_font_embedding(tmpl)
    assert not valid
    assert "invalid or missing os/2 table" in err.lower()


def test_v3_03_embed_multiple_font_families_and_clean_content_types(tmp_path):
    """Embedding with distinct body and heading font families embeds both families; stripping cleans [Content_Types].xml."""
    import yaml
    custom_dir = tmp_path / "multi_font_tmpl"
    shutil.copytree("templates/purple_book", custom_dir)

    # Create dummy Sahel font files by copying Vazirmatn files (which are valid TTF with OFL)
    fonts_dir = custom_dir / "fonts"
    shutil.copy(fonts_dir / "Vazirmatn-Regular.ttf", fonts_dir / "Sahel-Regular.ttf")
    shutil.copy(fonts_dir / "Vazirmatn-Bold.ttf", fonts_dir / "Sahel-Bold.ttf")

    cfg_path = custom_dir / "config.yaml"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    cfg["font_files"]["Sahel-Regular"] = "fonts/Sahel-Regular.ttf"
    cfg["font_files"]["Sahel-Bold"] = "fonts/Sahel-Bold.ttf"
    cfg_path.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")

    tmpl = Template.load(custom_dir)
    opts = GeneratorOptions(font_family="Vazirmatn", heading_font="Sahel", embed_fonts=True)

    out = tmp_path / "multi_embedded.docx"
    convert_markdown_to_docx(
        content="# تیتر با فونت ویژه\n\nمتن بدنه با فونت پیش‌فرض.\n",
        output_path=out,
        template=tmpl,
        options=opts,
        overwrite=True,
    )

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        # Both families embedded: 4 .odttf parts
        assert "word/fonts/font1.odttf" in names
        assert "word/fonts/font2.odttf" in names
        assert "word/fonts/font3.odttf" in names
        assert "word/fonts/font4.odttf" in names

        ft_xml = z.read("word/fontTable.xml").decode("utf-8")
        assert '<w:font w:name="Vazirmatn">' in ft_xml
        assert '<w:font w:name="Sahel">' in ft_xml

    ok, issues = run_package_oracle(out)
    assert ok, f"Package oracle reported issues: {issues}"

    # Now strip embedded fonts and verify [Content_Types].xml cleaned
    from md_to_docx.fonts_embed import strip_embedded_fonts
    strip_embedded_fonts(out)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert not any(n.startswith("word/fonts/") for n in names)
        ct_xml = z.read("[Content_Types].xml").decode("utf-8")
        assert 'Extension="odttf"' not in ct_xml


def test_v3_03_missing_heading_font_when_embedding_raises_error(tmp_path):
    """When embed_fonts=True, if heading_font files are missing, ConvertError is raised (Section 4.1)."""
    tmpl = Template.load("purple_book")
    opts = GeneratorOptions(font_family="Vazirmatn", heading_font="NonExistentHeadingFont", embed_fonts=True)
    out = tmp_path / "missing_heading_fail.docx"

    with pytest.raises(ConvertError, match="Font embedding failed for 'NonExistentHeadingFont'"):
        convert_markdown_to_docx(
            content="# تیتر\n\nمتن.\n",
            output_path=out,
            template=tmpl,
            options=opts,
            overwrite=True,
        )


def test_v3_04_package_oracle_detects_orphan_font(tmp_path):
    """Package oracle flags orphan font files in word/fonts/."""
    # Create valid embedded document
    out = tmp_path / "orphan_test.docx"
    convert_markdown_to_docx(
        content="# تست\n\nمتن.\n",
        output_path=out,
        template="purple_book",
        embed_fonts=True,
        overwrite=True,
    )

    # Inject an orphan font file into word/fonts/orphan.odttf
    orphan_doc = tmp_path / "with_orphan.docx"
    with zipfile.ZipFile(out, "r") as zin, zipfile.ZipFile(orphan_doc, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/fonts/orphan.odttf", b"fake font bytes")

    ok, issues = run_package_oracle(orphan_doc)
    assert not ok
    assert any("Orphan font part 'word/fonts/orphan.odttf'" in issue for issue in issues)


def test_v3_cli_prints_effective_direction_not_auto(tmp_path):
    """CLI prints effective direction (rtl or ltr), not unresolved 'direction: auto'."""
    runner = CliRunner()
    in_file = tmp_path / "persian.md"
    in_file.write_text("# عنوان فارسی\n\nاین یک متن فارسی است.\n", encoding="utf-8")
    out_file = tmp_path / "cli_persian.docx"

    result = runner.invoke(
        main,
        ["convert", str(in_file), "-o", str(out_file), "--template", "purple_book", "-f"],
    )
    assert result.exit_code == 0
    assert "direction: rtl" in result.output
    assert "direction: auto" not in result.output


