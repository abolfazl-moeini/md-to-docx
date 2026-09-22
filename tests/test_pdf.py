"""Comprehensive tests for PDF engine, pipeline integration, and CLI."""

from __future__ import annotations

import os
import shutil
import signal
import subprocess
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from click.testing import CliRunner

from md_to_docx import (
    convert_docx_to_pdf,
    convert_markdown_to_pdf,
    find_soffice_binary,
    is_valid_pdf,
    LibreOfficeNotFoundError,
)
from md_to_docx.cli import main
from md_to_docx.mermaid import ConvertError


# ---------------------------------------------------------------------------
# 1. Discovery tests (find_soffice_binary)
# ---------------------------------------------------------------------------

def test_find_soffice_binary_via_env_var(tmp_path, monkeypatch):
    dummy_bin = tmp_path / "soffice"
    dummy_bin.write_text("#!/bin/sh\nexit 0\n")
    dummy_bin.chmod(0o755)

    monkeypatch.setenv("MD2DOCX_SOFFICE", str(dummy_bin))
    found = find_soffice_binary()
    assert found == dummy_bin.resolve()


def test_find_soffice_binary_env_var_invalid_does_not_fall_back(monkeypatch):
    monkeypatch.setenv("MD2DOCX_SOFFICE", "/non/existent/soffice_path")
    monkeypatch.setattr("shutil.which", lambda name: "/usr/local/bin/soffice" if name == "soffice" else None)
    monkeypatch.setattr("pathlib.Path.is_file", lambda self: str(self) == "/usr/local/bin/soffice")

    assert find_soffice_binary() is None


def test_convert_docx_to_pdf_invalid_env_mentions_override(tmp_path, monkeypatch):
    monkeypatch.setenv("MD2DOCX_SOFFICE", "/non/existent/soffice_path")
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: None)
    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    with pytest.raises(LibreOfficeNotFoundError, match="is set to"):
        convert_docx_to_pdf(in_docx, tmp_path / "test.pdf")


def test_find_soffice_binary_via_which(monkeypatch):
    monkeypatch.delenv("MD2DOCX_SOFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda name: "/bin/libreoffice" if name == "libreoffice" else None)
    monkeypatch.setattr("pathlib.Path.is_file", lambda self: str(self) == "/bin/libreoffice")

    found = find_soffice_binary()
    assert found == Path("/bin/libreoffice").resolve()


def test_find_soffice_binary_standard_paths(monkeypatch, tmp_path):
    monkeypatch.delenv("MD2DOCX_SOFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _name: None)

    mac_app = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    monkeypatch.setattr("pathlib.Path.is_file", lambda self: self == mac_app)
    monkeypatch.setattr("os.access", lambda path, mode: True)

    found = find_soffice_binary()
    assert found == mac_app.resolve()


def test_find_soffice_binary_not_found(monkeypatch):
    monkeypatch.delenv("MD2DOCX_SOFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _name: None)
    monkeypatch.setattr("pathlib.Path.is_file", lambda self: False)

    assert find_soffice_binary() is None


# ---------------------------------------------------------------------------
# 2. PDF Validation tests (is_valid_pdf)
# ---------------------------------------------------------------------------

def test_is_valid_pdf_nonexistent_file(tmp_path):
    assert not is_valid_pdf(tmp_path / "nonexistent.pdf")


def test_is_valid_pdf_directory(tmp_path):
    sub = tmp_path / "sub"
    sub.mkdir()
    assert not is_valid_pdf(sub)


def test_is_valid_pdf_too_small(tmp_path):
    small = tmp_path / "small.pdf"
    small.write_bytes(b"%PDF-1.4 header\n%%EOF")
    # Default min_size is 1000
    assert not is_valid_pdf(small, min_size=1000)
    # But valid when min_size=10
    assert is_valid_pdf(small, min_size=10)


def test_is_valid_pdf_corrupted_header(tmp_path):
    corrupt = tmp_path / "corrupt.pdf"
    corrupt.write_bytes(b"NOT_A_PDF" + b"\x00" * 1200 + b"%%EOF")
    assert not is_valid_pdf(corrupt)


def test_is_valid_pdf_missing_eof_marker(tmp_path):
    no_eof = tmp_path / "no_eof.pdf"
    no_eof.write_bytes(b"%PDF-1.5\n" + b"some stream content" * 100)
    assert not is_valid_pdf(no_eof)


def test_is_valid_pdf_valid_file(tmp_path):
    valid = tmp_path / "valid.pdf"
    valid.write_bytes(b"%PDF-1.5\n" + b"body content padding " * 60 + b"\n%%EOF\n")
    assert valid.stat().st_size >= 1000
    assert is_valid_pdf(valid)


# ---------------------------------------------------------------------------
# 3. Adapter tests (convert_docx_to_pdf)
# ---------------------------------------------------------------------------

def test_convert_docx_to_pdf_missing_soffice_raises_error(monkeypatch, tmp_path):
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: None)
    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "test.pdf"

    with pytest.raises(LibreOfficeNotFoundError) as exc_info:
        convert_docx_to_pdf(in_docx, out_pdf)
    assert "PDF conversion requires LibreOffice" in str(exc_info.value)
    assert "brew install --cask libreoffice" in str(exc_info.value)


def test_convert_docx_to_pdf_missing_input_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        convert_docx_to_pdf(tmp_path / "missing.docx", tmp_path / "test.pdf")


def test_convert_docx_to_pdf_invalid_output_extension(tmp_path):
    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    with pytest.raises(ConvertError, match="must have a .pdf extension"):
        convert_docx_to_pdf(in_docx, tmp_path / "other.docx")


def test_convert_docx_to_pdf_existing_output_without_overwrite(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "test.pdf"
    out_pdf.write_bytes(b"existing")

    with pytest.raises(ConvertError, match="already exists"):
        convert_docx_to_pdf(in_docx, out_pdf, overwrite=False)


def test_convert_docx_to_pdf_mocked_success(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "sample.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "sample.pdf"

    cleaned_dirs = []
    original_rmtree = shutil.rmtree

    def mock_rmtree(path, *args, **kwargs):
        cleaned_dirs.append(str(path))
        original_rmtree(path, *args, **kwargs)

    monkeypatch.setattr("shutil.rmtree", mock_rmtree)

    def mock_popen(cmd, *args, **kwargs):
        # Extract --outdir
        outdir_idx = cmd.index("--outdir")
        stage_dir = Path(cmd[outdir_idx + 1])
        # Produce valid PDF in stage_dir
        pdf_file = stage_dir / "sample.pdf"
        pdf_file.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")

        mock_proc = MagicMock()
        mock_proc.communicate.return_value = ("Converted", "")
        mock_proc.returncode = 0
        mock_proc.poll.return_value = 0
        return mock_proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)

    res = convert_docx_to_pdf(in_docx, out_pdf)
    assert res == out_pdf
    assert out_pdf.exists()
    assert is_valid_pdf(out_pdf)
    # Profile dir and stage dir cleaned up
    assert any(".lo_prof_" in d for d in cleaned_dirs)
    assert any(".lo_stage_" in d for d in cleaned_dirs)


def test_convert_docx_to_pdf_timeout_kills_process_group(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "test.pdf"

    killed_pgid = []
    killed_signals = []

    mock_proc = MagicMock()
    mock_proc.pid = 9999
    mock_proc.poll.return_value = None
    mock_proc.communicate.side_effect = subprocess.TimeoutExpired(cmd=["soffice"], timeout=5)

    monkeypatch.setattr("subprocess.Popen", lambda *a, **kw: mock_proc)
    monkeypatch.setattr("os.getpgid", lambda pid: pid)
    monkeypatch.setattr(
        "os.killpg",
        lambda pgid, sig: (killed_pgid.append(pgid), killed_signals.append(sig)),
    )

    cleaned_dirs = []
    original_rmtree = shutil.rmtree

    def mock_rmtree(path, *args, **kwargs):
        cleaned_dirs.append(str(path))
        original_rmtree(path, *args, **kwargs)

    monkeypatch.setattr("shutil.rmtree", mock_rmtree)

    with pytest.raises(ConvertError, match="timed out"):
        convert_docx_to_pdf(in_docx, out_pdf, timeout=5)

    assert 9999 in killed_pgid
    assert signal.SIGKILL in killed_signals
    # Cleanup occurred
    assert any(".lo_prof_" in d for d in cleaned_dirs)


def test_convert_docx_to_pdf_failure_exit_code(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "test.pdf"

    mock_proc = MagicMock()
    mock_proc.returncode = 1
    mock_proc.communicate.return_value = ("", "Fatal: segmentation fault in soffice.bin")

    monkeypatch.setattr("subprocess.Popen", lambda *a, **kw: mock_proc)

    with pytest.raises(ConvertError, match="LibreOffice conversion failed.*segmentation fault"):
        convert_docx_to_pdf(in_docx, out_pdf)


# ---------------------------------------------------------------------------
# 4. Pipeline tests (convert_markdown_to_pdf)
# ---------------------------------------------------------------------------

def test_convert_markdown_to_pdf_rejects_non_pdf_output(tmp_path):
    with pytest.raises(ConvertError, match="must have a .pdf extension"):
        convert_markdown_to_pdf(content="# Test", output_path=tmp_path / "out.docx")


def test_convert_markdown_to_pdf_missing_libreoffice_raises(monkeypatch, tmp_path):
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: None)
    with pytest.raises(LibreOfficeNotFoundError):
        convert_markdown_to_pdf(content="# Test", output_path=tmp_path / "out.pdf")


def test_convert_markdown_to_pdf_media_hygiene_and_keep_docx(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    # Mock convert_docx_to_pdf to write a valid PDF file
    def mock_convert_docx(docx_path, output_pdf_path, **kwargs):
        p = Path(output_pdf_path)
        p.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        return p

    monkeypatch.setattr("md_to_docx.pipeline.convert_docx_to_pdf", mock_convert_docx)

    # Markdown with simulated mermaid diagram
    md_content = "# Title\n\n```mermaid\ngraph TD;\nA-->B;\n```\n"

    # Mock render_mermaid_fn to write a diagram PNG
    from PIL import Image as PILImage
    def dummy_mermaid_fn(code, out_png, template, options=None):
        im = PILImage.new("RGB", (10, 10), color="blue")
        im.save(out_png, format="PNG")
        return True

    # CASE A: keep_docx=False (default) -> no media directory should be published
    out_pdf_a = tmp_path / "doc_a.pdf"
    media_dir_a = tmp_path / "doc_a_media"
    res_a = convert_markdown_to_pdf(
        content=md_content,
        output_path=out_pdf_a,
        render_mermaid_fn=dummy_mermaid_fn,
        keep_docx=False,
    )
    assert res_a == out_pdf_a
    assert out_pdf_a.exists()
    assert not media_dir_a.exists(), "Media directory must NOT be published when keep_docx=False"
    assert not (tmp_path / "doc_a.docx").exists()

    # CASE B: keep_docx=True -> intermediate docx AND media directory are published
    out_pdf_b = tmp_path / "doc_b.pdf"
    media_dir_b = tmp_path / "doc_b_media"
    res_b = convert_markdown_to_pdf(
        content=md_content,
        output_path=out_pdf_b,
        render_mermaid_fn=dummy_mermaid_fn,
        keep_docx=True,
    )
    assert res_b == out_pdf_b
    assert out_pdf_b.exists()
    assert (tmp_path / "doc_b.docx").exists()
    assert media_dir_b.exists()
    assert list(media_dir_b.glob("diagram_*.png"))


def test_convert_markdown_to_pdf_rollback_on_failure(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    out_pdf = tmp_path / "rollback.pdf"
    out_pdf.write_bytes(b"%PDF-1.4\nORIGINAL CONTENT" + b" " * 1000 + b"\n%%EOF\n")

    # Mock convert_docx_to_pdf to raise an error
    def failing_convert(*args, **kwargs):
        raise ConvertError("LibreOffice internal failure")

    monkeypatch.setattr("md_to_docx.pipeline.convert_docx_to_pdf", failing_convert)

    with pytest.raises(ConvertError, match="LibreOffice internal failure"):
        convert_markdown_to_pdf(
            content="# New title",
            output_path=out_pdf,
            overwrite=True,
        )

    # Ensure original content was preserved / restored
    assert out_pdf.exists()
    assert b"ORIGINAL CONTENT" in out_pdf.read_bytes()


# ---------------------------------------------------------------------------
# 5. CLI tests (convert -o *.pdf, to-pdf)
# ---------------------------------------------------------------------------

@pytest.fixture
def runner():
    return CliRunner()


def test_cli_convert_to_pdf_missing_libreoffice_clean_error(runner, tmp_path, monkeypatch):
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: None)
    in_md = tmp_path / "doc.md"
    in_md.write_text("# Test", encoding="utf-8")
    out_pdf = tmp_path / "doc.pdf"

    result = runner.invoke(main, ["convert", str(in_md), "-o", str(out_pdf)])
    assert result.exit_code == 1
    assert "PDF conversion requires LibreOffice ('soffice')." in result.output
    assert "Traceback" not in result.output


def test_cli_convert_to_pdf_success_mocked(runner, tmp_path, mocker):
    in_md = tmp_path / "doc.md"
    in_md.write_text("# عنوان تست", encoding="utf-8")
    out_pdf = tmp_path / "doc.pdf"

    mocker.patch("md_to_docx.cli.convert_markdown_to_pdf", return_value=out_pdf)

    result = runner.invoke(main, ["convert", str(in_md), "-o", str(out_pdf)])
    assert result.exit_code == 0
    assert "Success: Generated PDF" in result.output


def test_cli_convert_rejects_unsupported_extensions(runner, tmp_path):
    in_md = tmp_path / "doc.md"
    in_md.write_text("# Test", encoding="utf-8")

    result = runner.invoke(main, ["convert", str(in_md), "-o", str(tmp_path / "out.html")])
    assert result.exit_code == 2
    assert "must have a .docx or .pdf extension" in result.output


def test_cli_to_pdf_validations(runner, tmp_path):
    # Reject non-docx input
    txt_file = tmp_path / "file.txt"
    txt_file.write_text("not docx", encoding="utf-8")
    result = runner.invoke(main, ["to-pdf", str(txt_file)])
    assert result.exit_code == 2
    assert "must have a .docx extension" in result.output

    # Reject .doc input
    doc_file = tmp_path / "file.doc"
    doc_file.write_bytes(b"old doc")
    result = runner.invoke(main, ["to-pdf", str(doc_file)])
    assert result.exit_code == 2
    assert "Word 97-2003 .doc is not supported" in result.output

    # Reject nonexistent input
    result = runner.invoke(main, ["to-pdf", str(tmp_path / "nonexistent.docx")])
    assert result.exit_code == 2
    assert "does not exist" in result.output

    # Reject non-pdf output
    valid_docx = tmp_path / "valid.docx"
    valid_docx.write_bytes(b"PK\x03\x04")
    result = runner.invoke(main, ["to-pdf", str(valid_docx), "-o", str(tmp_path / "out.txt")])
    assert result.exit_code == 2
    assert "must have a .pdf extension" in result.output


def test_cli_to_pdf_missing_libreoffice_clean_error(runner, tmp_path, monkeypatch):
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: None)
    in_docx = tmp_path / "sample.docx"
    in_docx.write_bytes(b"PK\x03\x04")

    result = runner.invoke(main, ["to-pdf", str(in_docx)])
    assert result.exit_code == 1
    assert "PDF conversion requires LibreOffice ('soffice')." in result.output
    assert "Traceback" not in result.output


def test_cli_to_pdf_success_mocked(runner, tmp_path, mocker):
    in_docx = tmp_path / "doc.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "doc.pdf"

    mocker.patch("md_to_docx.cli.convert_docx_to_pdf", return_value=out_pdf)

    result = runner.invoke(main, ["to-pdf", str(in_docx), "-o", str(out_pdf)])
    assert result.exit_code == 0
    assert "Success: Generated PDF" in result.output


def test_cli_to_pdf_default_template_resolves_vazirmatn(runner, tmp_path, mocker):
    in_docx = tmp_path / "doc.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "doc.pdf"
    mock = mocker.patch("md_to_docx.cli.convert_docx_to_pdf", return_value=out_pdf)

    result = runner.invoke(main, ["to-pdf", str(in_docx), "-o", str(out_pdf), "-f"])
    assert result.exit_code == 0, result.output
    font_dirs = mock.call_args.kwargs.get("font_dirs") or []
    assert any((Path(d) / "Vazirmatn-Regular.ttf").is_file() for d in font_dirs)


def test_cli_to_pdf_template_resolves_font_dir(runner, tmp_path, mocker):
    """--template must point LibreOffice at the template font folder, not a relative fonts/ path."""
    in_docx = tmp_path / "doc.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "doc.pdf"
    mock = mocker.patch("md_to_docx.cli.convert_docx_to_pdf", return_value=out_pdf)

    result = runner.invoke(
        main,
        ["to-pdf", str(in_docx), "-o", str(out_pdf), "--template", "purple_book", "-f"],
    )
    assert result.exit_code == 0, result.output
    font_dirs = mock.call_args.kwargs.get("font_dirs") or []
    assert any((Path(d) / "Vazirmatn-Regular.ttf").is_file() for d in font_dirs)


def test_cli_convert_stdin_to_pdf_success_mocked(runner, tmp_path, mocker):
    out_pdf = tmp_path / "from_stdin.pdf"
    mocker.patch("md_to_docx.cli.convert_markdown_to_pdf", return_value=out_pdf)

    result = runner.invoke(main, ["convert", "-", "-o", str(out_pdf)], input="# From stdin\n\nPersian text")
    assert result.exit_code == 0
    assert "Success: Generated PDF" in result.output


def test_cli_convert_to_pdf_flags_forwarded(runner, tmp_path, mocker):
    in_md = tmp_path / "doc.md"
    in_md.write_text("# Test", encoding="utf-8")
    out_pdf = tmp_path / "doc.pdf"

    mock_fn = mocker.patch("md_to_docx.cli.convert_markdown_to_pdf", return_value=out_pdf)

    result = runner.invoke(main, ["convert", str(in_md), "-o", str(out_pdf), "--keep-docx", "--pdf-timeout", "45"])
    assert result.exit_code == 0
    assert mock_fn.called
    kwargs = mock_fn.call_args.kwargs
    assert kwargs.get("keep_docx") is True
    assert kwargs.get("pdf_timeout") == 45


def test_convert_docx_to_pdf_rejects_mismatched_stem_pdf(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "sample.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "sample.pdf"

    def mock_popen(cmd, *args, **kwargs):
        stage_dir = Path(cmd[cmd.index("--outdir") + 1])
        (stage_dir / "unrelated.pdf").write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        mock_proc = MagicMock()
        mock_proc.communicate.return_value = ("Done", "")
        mock_proc.returncode = 0
        mock_proc.poll.return_value = 0
        return mock_proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)

    with pytest.raises(ConvertError, match="no PDF output was found"):
        convert_docx_to_pdf(in_docx, out_pdf)


def test_convert_docx_to_pdf_produces_corrupt_pdf_fails(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "sample.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "sample.pdf"

    def mock_popen(cmd, *args, **kwargs):
        outdir_idx = cmd.index("--outdir")
        stage_dir = Path(cmd[outdir_idx + 1])
        # Produce corrupt PDF (not starting with %PDF-)
        pdf_file = stage_dir / "sample.pdf"
        pdf_file.write_bytes(b"CORRUPT_BYTES" * 100)

        mock_proc = MagicMock()
        mock_proc.communicate.return_value = ("Done", "")
        mock_proc.returncode = 0
        mock_proc.poll.return_value = 0
        return mock_proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)

    with pytest.raises(ConvertError, match="invalid, empty, or corrupt"):
        convert_docx_to_pdf(in_docx, out_pdf)


def test_cli_to_pdf_already_exists_without_overwrite(runner, tmp_path):
    in_docx = tmp_path / "doc.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "doc.pdf"
    out_pdf.write_bytes(b"existing")

    result = runner.invoke(main, ["to-pdf", str(in_docx), "-o", str(out_pdf)])
    assert result.exit_code == 2
    assert "already exists" in result.output


def test_cli_convert_keep_docx_already_exists_exits_code_2(runner, tmp_path):
    in_md = tmp_path / "test.md"
    in_md.write_text("# Title", encoding="utf-8")
    out_pdf = tmp_path / "test.pdf"
    existing_docx = tmp_path / "test.docx"
    existing_docx.write_bytes(b"existing docx")

    result = runner.invoke(main, ["convert", str(in_md), "-o", str(out_pdf), "--keep-docx"])
    assert result.exit_code == 2
    assert "Intermediate DOCX file" in result.output
    assert "already exists" in result.output


def test_cli_convert_stdin_keep_docx_already_exists_exits_code_2(runner, tmp_path):
    out_pdf = tmp_path / "stdin_test.pdf"
    existing_docx = tmp_path / "stdin_test.docx"
    existing_docx.write_bytes(b"existing docx")

    result = runner.invoke(
        main,
        ["convert", "-", "-o", str(out_pdf), "--keep-docx"],
        input="# From stdin",
    )
    assert result.exit_code == 2
    assert "Intermediate DOCX file" in result.output
    assert "already exists" in result.output


def test_cli_convert_invalid_pdf_timeout_exits_code_2(runner, tmp_path):
    in_md = tmp_path / "doc.md"
    in_md.write_text("# Test", encoding="utf-8")
    result = runner.invoke(main, ["convert", str(in_md), "-o", str(tmp_path / "out.pdf"), "--pdf-timeout", "0"])
    assert result.exit_code == 2
    assert "not in the range" in result.output or "Error" in result.output


def test_cli_to_pdf_invalid_timeout_exits_code_2(runner, tmp_path):
    in_docx = tmp_path / "doc.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    result = runner.invoke(main, ["to-pdf", str(in_docx), "--pdf-timeout", "-5"])
    assert result.exit_code == 2
    assert "not in the range" in result.output or "Error" in result.output


def test_convert_markdown_to_pdf_fails_fast_on_existing_output(tmp_path, monkeypatch):
    out_pdf = tmp_path / "existing.pdf"
    out_pdf.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")

    pandoc_called = False
    def mock_pandoc(*args, **kwargs):
        nonlocal pandoc_called
        pandoc_called = True
        return {}

    monkeypatch.setattr("md_to_docx.pipeline.run_pandoc_ast", mock_pandoc)

    with pytest.raises(ConvertError, match="already exists"):
        convert_markdown_to_pdf(content="# Test", output_path=out_pdf, overwrite=False)

    assert not pandoc_called, "convert_markdown_to_pdf must fail fast before AST parsing"


def test_convert_markdown_to_pdf_passes_font_dirs_from_template(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    received_font_dirs = []
    def mock_convert_docx(docx_path, output_pdf_path, font_dirs=None, **kwargs):
        if font_dirs:
            received_font_dirs.extend(font_dirs)
        p = Path(output_pdf_path)
        p.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        return p

    monkeypatch.setattr("md_to_docx.pdf.convert_docx_to_pdf", mock_convert_docx)

    out_pdf = tmp_path / "font_test.pdf"
    convert_markdown_to_pdf(content="# Title", output_path=out_pdf, template="purple_book")

    assert len(received_font_dirs) > 0
    assert any("fonts" in str(p) for p in received_font_dirs)


def test_convert_docx_to_pdf_generates_fontconfig_conf(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    fake_font_dir = tmp_path / "fonts"
    fake_font_dir.mkdir()
    (fake_font_dir / "Vazirmatn.ttf").write_bytes(b"dummy font")

    captured_env = {}
    def mock_popen(cmd, env=None, *args, **kwargs):
        if env:
            captured_env.update(env)
        fc_text = ""
        if env and "FONTCONFIG_FILE" in env:
            fc_file = Path(env["FONTCONFIG_FILE"])
            if fc_file.exists():
                fc_text = fc_file.read_text(encoding="utf-8")
        captured_env["fc_text"] = fc_text

        outdir_idx = cmd.index("--outdir")
        stage_dir = Path(cmd[outdir_idx + 1])
        pdf_file = stage_dir / "test.pdf"
        pdf_file.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        proc = MagicMock()
        proc.returncode = 0
        proc.communicate.return_value = ("", "")
        proc.poll.return_value = 0
        return proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)

    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "test.pdf"

    convert_docx_to_pdf(in_docx, out_pdf, font_dirs=[fake_font_dir])

    assert "SAL_FONTPATH" in captured_env
    assert str(fake_font_dir.resolve()) in captured_env["SAL_FONTPATH"]
    assert "FONTCONFIG_FILE" in captured_env
    fc_path = Path(captured_env["FONTCONFIG_FILE"])
    assert fc_path.name == "fonts.conf"
    # Content was verified while active
    assert str(fake_font_dir.resolve()) in captured_env["fc_text"]
    # Profile dir was cleaned up after completion
    assert not fc_path.exists()


def test_convert_docx_to_pdf_escapes_font_dir_in_fontconfig(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    fake_font_dir = tmp_path / "fonts & <extra>"
    fake_font_dir.mkdir()
    captured = {}

    def mock_popen(cmd, env=None, *args, **kwargs):
        fc_text = Path(env["FONTCONFIG_FILE"]).read_text(encoding="utf-8") if env and env.get("FONTCONFIG_FILE") else ""
        captured["fc_text"] = fc_text
        stage_dir = Path(cmd[cmd.index("--outdir") + 1])
        (stage_dir / "test.pdf").write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        proc = MagicMock()
        proc.returncode = 0
        proc.communicate.return_value = ("", "")
        proc.poll.return_value = 0
        return proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)
    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    convert_docx_to_pdf(in_docx, tmp_path / "test.pdf", font_dirs=[fake_font_dir])
    assert "&amp;" in captured["fc_text"]
    assert "&lt;extra&gt;" in captured["fc_text"]
    assert "<extra>" not in captured["fc_text"]


def test_convert_docx_to_pdf_identical_input_output_raises(tmp_path):
    same_path = tmp_path / "same.docx"
    same_path.write_bytes(b"PK\x03\x04")
    with pytest.raises(ValueError, match="cannot be identical"):
        convert_docx_to_pdf(same_path, same_path)


def test_convert_docx_to_pdf_non_positive_timeout_raises(tmp_path):
    in_docx = tmp_path / "test.docx"
    in_docx.write_bytes(b"PK\x03\x04")
    out_pdf = tmp_path / "test.pdf"
    with pytest.raises(ValueError, match="timeout must be positive"):
        convert_docx_to_pdf(in_docx, out_pdf, timeout=0)
    with pytest.raises(ValueError, match="timeout must be positive"):
        convert_docx_to_pdf(in_docx, out_pdf, timeout=-10)


def test_convert_markdown_to_pdf_non_positive_timeout_raises(tmp_path):
    with pytest.raises(ValueError, match="pdf_timeout must be positive"):
        convert_markdown_to_pdf(content="# Test", output_path=tmp_path / "out.pdf", pdf_timeout=0)


def test_kill_process_tree_windows(monkeypatch):
    from md_to_docx.pdf import _kill_process_tree

    monkeypatch.setattr("os.name", "nt")
    ran_taskkill = []
    def mock_run(cmd, *args, **kwargs):
        ran_taskkill.append(cmd)
        return MagicMock(returncode=0)

    monkeypatch.setattr("subprocess.run", mock_run)

    mock_proc = MagicMock()
    mock_proc.pid = 4321
    mock_proc.poll.return_value = None

    _kill_process_tree(mock_proc)
    assert any("taskkill" in cmd[0] and "4321" in cmd for cmd in ran_taskkill)


def test_find_soffice_binary_non_executable_posix(tmp_path, monkeypatch):
    monkeypatch.setattr("os.name", "posix")
    dummy = tmp_path / "non_exec_soffice"
    dummy.write_text("#!/bin/sh\n")
    dummy.chmod(0o644)  # Not executable

    monkeypatch.setenv("MD2DOCX_SOFFICE", str(dummy))
    monkeypatch.setattr("shutil.which", lambda _n: None)
    monkeypatch.setattr("pathlib.Path.is_file", lambda self: str(self) == str(dummy))

    assert find_soffice_binary() is None


def test_convert_docx_to_pdf_rejects_non_docx_input(tmp_path):
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("not a docx", encoding="utf-8")
    with pytest.raises(ConvertError, match="must have a .docx extension"):
        convert_docx_to_pdf(txt_file, tmp_path / "out.pdf")

    doc_file = tmp_path / "legacy.doc"
    doc_file.write_bytes(b"old word")
    with pytest.raises(ConvertError, match="Word 97-2003 .doc is not supported"):
        convert_docx_to_pdf(doc_file, tmp_path / "out.pdf")


def test_is_valid_pdf_rejects_xref_without_eof(tmp_path):
    truncated = tmp_path / "truncated.pdf"
    truncated.write_bytes(b"%PDF-1.5\n" + b"A" * 1100 + b" startxref 12345 xref")
    assert not is_valid_pdf(truncated)


def test_convert_docx_to_pdf_leaves_publish_lock(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    in_docx = tmp_path / "locked.docx"
    in_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "locked.pdf"

    def mock_popen(cmd, *args, **kwargs):
        stage_dir = Path(cmd[cmd.index("--outdir") + 1])
        (stage_dir / "locked.pdf").write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        mock_proc = MagicMock()
        mock_proc.communicate.return_value = ("Converted", "")
        mock_proc.returncode = 0
        mock_proc.poll.return_value = 0
        return mock_proc

    monkeypatch.setattr("subprocess.Popen", mock_popen)

    res = convert_docx_to_pdf(in_docx, out_pdf)
    assert res == out_pdf
    # Lock file is intentionally retained so concurrent publishers share the inode.
    assert (tmp_path / ".locked.pdf.publish.lock").exists()


def test_convert_markdown_to_pdf_report_keys(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    def mock_convert_docx(docx_path, output_pdf_path, **kwargs):
        p = Path(output_pdf_path)
        p.write_bytes(b"%PDF-1.4\n" + b"X" * 1200 + b"\n%%EOF\n")
        return p

    monkeypatch.setattr("md_to_docx.pipeline.convert_docx_to_pdf", mock_convert_docx)

    out_pdf = tmp_path / "rep.pdf"
    report: dict = {}
    convert_markdown_to_pdf(
        content="# Title\n",
        output_path=out_pdf,
        report=report,
        keep_docx=True,
        pdf_timeout=45,
    )
    assert report["pdf_engine"] == "libreoffice"
    assert report["pdf_timeout"] == 45
    assert report["keep_docx"] is True
    assert report["intermediate_docx"] == str(tmp_path / "rep.docx")
    assert "effective_body_font" in report


# ---------------------------------------------------------------------------
# 6. Live integration test (marked with pytest.mark.pdf)
# ---------------------------------------------------------------------------

@pytest.mark.pdf
@pytest.mark.integration
def test_live_pdf_conversion_e2e(tmp_path):
    """End-to-end live conversion if LibreOffice is present on the machine."""
    soffice = find_soffice_binary()
    if not soffice:
        if os.environ.get("MD2DOCX_REQUIRE_EXTERNAL") == "1":
            pytest.fail("LibreOffice ('soffice') binary is required by MD2DOCX_REQUIRE_EXTERNAL=1 but not found.")
        else:
            pytest.skip("LibreOffice ('soffice') not found on system; skipping live integration test.")

    md_content = """# گزارش آزمایشی چیدمان فارسی

این یک متن آزمایشی به زبان فارسی است که به پی‌دی‌اف تبدیل می‌شود.

- آیتم اول
- آیتم دوم

| ردیف | نام | وضعیت |
| --- | --- | --- |
| ۱ | آزمون A | فعال |
| ۲ | آزمون B | غیرفعال |
"""
    out_pdf = tmp_path / "output_live.pdf"
    result = convert_markdown_to_pdf(
        content=md_content,
        output_path=out_pdf,
        template="purple_book",
        overwrite=True,
    )
    assert result == out_pdf
    assert out_pdf.exists()
    assert is_valid_pdf(out_pdf)


def test_convert_docx_to_pdf_identical_docx_paths_raises(tmp_path):
    same_docx = tmp_path / "same.docx"
    same_docx.write_bytes(b"PK\x03\x04dummy")
    with pytest.raises(ValueError, match="cannot be identical"):
        convert_docx_to_pdf(same_docx, same_docx)


def test_convert_docx_to_pdf_directories_raise(tmp_path):
    sub = tmp_path / "sub"
    sub.mkdir()
    valid_docx = tmp_path / "valid.docx"
    valid_docx.write_bytes(b"PK\x03\x04dummy")
    out_pdf = tmp_path / "out.pdf"

    with pytest.raises(IsADirectoryError, match="is a directory"):
        convert_docx_to_pdf(sub, out_pdf)

    with pytest.raises(IsADirectoryError, match="is a directory"):
        convert_docx_to_pdf(valid_docx, sub)


def test_convert_markdown_to_pdf_missing_input_and_content_raises():
    with pytest.raises(ValueError, match="Either input_path or content must be provided"):
        convert_markdown_to_pdf(input_path=None, content=None)


def test_convert_markdown_to_pdf_invalid_options_type_raises(tmp_path):
    with pytest.raises(ValueError, match="Invalid options type"):
        convert_markdown_to_pdf(content="# Test", output_path=tmp_path / "out.pdf", options="invalid_str")


def test_convert_markdown_to_pdf_template_safety_guards(tmp_path, monkeypatch):
    dummy_soffice = tmp_path / "soffice"
    dummy_soffice.write_text("#!/bin/sh\n")
    dummy_soffice.chmod(0o755)
    monkeypatch.setattr("md_to_docx.pdf.find_soffice_binary", lambda: dummy_soffice)

    from md_to_docx.template import Template
    tmpl = Template.load("purple_book")

    # Output cannot be inside template directory
    with pytest.raises(ConvertError, match="cannot be inside the template directory"):
        convert_markdown_to_pdf(
            content="# Test",
            output_path=tmpl.dir_path / "leak.pdf",
            template=tmpl,
        )

    # Intermediate DOCX cannot overwrite template shell file
    custom_shell = tmp_path / "custom_shell.docx"
    custom_shell.write_bytes(b"PK\x03\x04")
    tmpl.shell_docx_path = custom_shell
    with pytest.raises(ConvertError, match="Intermediate DOCX path cannot overwrite template shell file"):
        convert_markdown_to_pdf(
            content="# Test",
            output_path=custom_shell.with_suffix(".pdf"),
            template=tmpl,
            keep_docx=True,
        )


def test_convert_markdown_to_pdf_size_limit_exceeded(tmp_path):
    from md_to_docx.pipeline import MAX_INPUT_SIZE_BYTES
    oversized = "a" * (MAX_INPUT_SIZE_BYTES + 10)
    with pytest.raises(ConvertError, match="exceeds maximum supported limit"):
        convert_markdown_to_pdf(content=oversized, output_path=tmp_path / "out.pdf")


def test_cli_to_pdf_additional_validations(runner, tmp_path):
    valid_docx = tmp_path / "input.docx"
    valid_docx.write_bytes(b"PK\x03\x04")

    # Identical input and output path
    res = runner.invoke(main, ["to-pdf", str(valid_docx), "-o", str(valid_docx)])
    assert res.exit_code == 2
    assert "Output path cannot be identical to input path" in res.output

    # Unsupported CLI flags should be rejected by click with exit code 2
    res_flags = runner.invoke(main, ["to-pdf", str(valid_docx), "--font", "Vazirmatn"])
    assert res_flags.exit_code == 2
    assert "no such option" in res_flags.output.lower() or "unrecognized" in res_flags.output.lower()


def test_package_metadata_and_exports():
    import md_to_docx
    assert md_to_docx.__version__ == "0.3.0"
    for symbol in md_to_docx.__all__:
        assert hasattr(md_to_docx, symbol), f"Missing public symbol '{symbol}' in md_to_docx"


def test_docx_is_rtl_detection(tmp_path):
    from md_to_docx.pdf import _docx_is_rtl
    from md_to_docx.oxml import set_doc_bidi, set_paragraph_bidi, set_table_bidi_visual
    from docx import Document

    # 1. Non-RTL docx
    doc_ltr = Document()
    doc_ltr.add_paragraph("Hello world")
    p_ltr = tmp_path / "ltr.docx"
    doc_ltr.save(str(p_ltr))
    assert _docx_is_rtl(p_ltr) is False

    # 2. Section-level RTL (the catalog Direction signal)
    doc_rtl = Document()
    doc_rtl.add_paragraph("سلام دنیا")
    set_doc_bidi(doc_rtl, bidi=True)
    p_rtl = tmp_path / "rtl.docx"
    doc_rtl.save(str(p_rtl))
    assert _docx_is_rtl(p_rtl) is True

    # 3. Explicit LTR section bidi=0 must not be treated as RTL
    doc_ltr_explicit = Document()
    doc_ltr_explicit.add_paragraph("Hello")
    set_doc_bidi(doc_ltr_explicit, bidi=False)
    p_ltr_explicit = tmp_path / "ltr_explicit.docx"
    doc_ltr_explicit.save(str(p_ltr_explicit))
    assert _docx_is_rtl(p_ltr_explicit) is False

    # 4. Paragraph w:bidi or table bidiVisual alone is not document RTL
    doc_para = Document()
    p = doc_para.add_paragraph("سلام")
    set_paragraph_bidi(p, bidi=True)
    tbl = doc_para.add_table(rows=1, cols=1)
    set_table_bidi_visual(tbl)
    p_para = tmp_path / "para_bidi.docx"
    doc_para.save(str(p_para))
    assert _docx_is_rtl(p_para) is False

    # 5. Converter LTR output must not look like RTL (sectPr bidi=0 still contains <w:bidi)
    from md_to_docx import convert_markdown_to_docx
    ltr_out = tmp_path / "converted_ltr.docx"
    convert_markdown_to_docx(
        content="# Hello\n\nAn English paragraph.\n",
        output_path=ltr_out,
        template="purple_book",
        direction="ltr",
        overwrite=True,
    )
    assert _docx_is_rtl(ltr_out) is False

    rtl_out = tmp_path / "converted_rtl.docx"
    convert_markdown_to_docx(
        content="# سلام\n\nیک پاراگراف فارسی.\n",
        output_path=rtl_out,
        template="purple_book",
        direction="rtl",
        overwrite=True,
    )
    assert _docx_is_rtl(rtl_out) is True

    # 6. Non-existent file
    assert _docx_is_rtl(tmp_path / "nonexistent.docx") is False


def test_apply_pdf_r2l_direction(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    from md_to_docx.pdf import _apply_pdf_r2l_direction

    # Create a minimal valid PDF
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    test_pdf = tmp_path / "test_catalog.pdf"
    with open(test_pdf, "wb") as f:
        writer.write(f)

    # Initially has no ViewerPreferences
    reader = pypdf.PdfReader(str(test_pdf))
    assert "/ViewerPreferences" not in reader.trailer["/Root"]

    # Apply R2L
    _apply_pdf_r2l_direction(test_pdf)

    # Verify ViewerPreferences << /Direction /R2L >> is injected
    reader_after = pypdf.PdfReader(str(test_pdf))
    root = reader_after.trailer["/Root"]
    assert "/ViewerPreferences" in root
    vp = root["/ViewerPreferences"]
    assert vp["/Direction"] == "/R2L"
    assert is_valid_pdf(test_pdf, min_size=32)


def test_apply_pdf_r2l_keeps_original_when_rewrite_is_invalid(tmp_path, monkeypatch):
    pypdf = pytest.importorskip("pypdf")
    from md_to_docx import pdf as pdf_mod

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    test_pdf = tmp_path / "keep_original.pdf"
    with open(test_pdf, "wb") as f:
        writer.write(f)
    original = test_pdf.read_bytes()

    def boom(*_args, **_kwargs):
        raise RuntimeError("pypdf rewrite failed")

    monkeypatch.setattr(pypdf.PdfWriter, "write", boom)
    pdf_mod._apply_pdf_r2l_direction(test_pdf)
    assert test_pdf.read_bytes() == original
    leftovers = list(tmp_path.glob(".keep_original.r2l_*.tmp"))
    assert leftovers == []

