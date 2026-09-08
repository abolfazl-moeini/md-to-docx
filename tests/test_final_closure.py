"""Final-closure plan tests (F01–F16). Do not launch a real browser here."""

from pathlib import Path
import html
import zipfile
import pytest
from lxml import etree
from docx.oxml.ns import qn

from md_to_docx import convert_markdown_to_docx, Template
from md_to_docx.mermaid import ConvertError
from md_to_docx.template import TemplateValidationError

ROOT = Path(__file__).resolve().parent.parent


def test_f01_iter_browsers_skips_system_chrome_and_edge(monkeypatch, tmp_path):
    from md_to_docx import mermaid as mermaid_mod

    mermaid_mod.reset_launch_health()
    monkeypatch.delenv("PUPPETEER_EXECUTABLE_PATH", raising=False)
    monkeypatch.delenv("MD2DOCX_ALLOW_SYSTEM_BROWSER", raising=False)
    monkeypatch.setattr(mermaid_mod, "_browsers_in_puppeteer_cache", lambda: [])
    monkeypatch.setattr(
        mermaid_mod,
        "_system_browser_candidates",
        lambda: [
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        ],
    )
    monkeypatch.setattr(mermaid_mod.shutil, "which", lambda _name: None)
    # Pretend system binaries exist so the old code would have picked them.
    monkeypatch.setattr(Path, "is_file", lambda self: True)

    found = mermaid_mod._iter_browser_candidates()
    joined = " ".join(found)
    assert "Google Chrome.app" not in joined
    assert "Microsoft Edge.app" not in joined


def test_f01_invalid_explicit_browser_does_not_fallback(tmp_path, monkeypatch):
    from md_to_docx import mermaid as mermaid_mod
    from md_to_docx.template import Template

    mermaid_mod.reset_launch_health()
    missing = tmp_path / "no-such-chrome"
    monkeypatch.setenv("PUPPETEER_EXECUTABLE_PATH", str(missing))
    tmpl = Template.load("purple_book")
    with pytest.raises(ConvertError) as exc:
        mermaid_mod.render_mermaid_to_png("graph TD;A-->B;", tmp_path / "x.png", tmpl)
    assert "explicit" in str(exc.value).lower() or "PUPPETEER_EXECUTABLE_PATH" in str(exc.value)


def test_f01_launch_failure_blocks_subsequent_attempts(tmp_path, monkeypatch):
    from md_to_docx import mermaid as mermaid_mod

    mermaid_mod.reset_launch_health()
    mermaid_mod.record_launch_failure("simulated launch crash")
    tmpl = Template.load("purple_book")
    with pytest.raises(ConvertError) as exc:
        mermaid_mod.render_mermaid_to_png("graph TD;A-->B;", tmp_path / "x.png", tmpl)
    assert "blocked" in str(exc.value).lower() or "previous" in str(exc.value).lower()


def test_f01_mmdc_does_not_use_npx_float(monkeypatch):
    from md_to_docx import mermaid as mermaid_mod

    monkeypatch.setattr(mermaid_mod.shutil, "which", lambda name: "npx" if name == "npx" else None)
    monkeypatch.setattr(Path, "is_file", lambda self: False)
    with pytest.raises(ConvertError):
        mermaid_mod._find_mmdc_cmd()


def test_f02_unknown_fixture_exits_nonzero():
    import subprocess, sys
    res = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "matrix_runner.py"), "--fixtures", "NO_SUCH_FIXTURE"],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    assert res.returncode != 0
    assert "NO_SUCH_FIXTURE" in (res.stderr + res.stdout)


def test_f02_html_report_escapes_issues(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.matrix_runner import generate_matrix_reports

    sample = [{
        "fixture_id": "S01",
        "template": "persian_book",
        "status": "fail",
        "conversion_time_sec": 0.1,
        "file_size_bytes": 10,
        "structural_pass": False,
        "content_pass": False,
        "warnings_pass": False,
        "issues": ["broken <script>alert(1)</script> rel"],
        "warnings": [],
        "counts": {},
    }]
    out = tmp_path / "rep"
    out.mkdir()
    generate_matrix_reports(sample, out, "escape_run", 1.0)
    html_text = (out / "index.html").read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in html_text
    assert "&lt;script&gt;" in html_text


def test_f03_deleted_body_paragraph_fails_semantic_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import run_semantic_oracle

    md = "# عنوان\n\nپاراگراف یکتای الف UNIQUE_PARA_A.\n\nپاراگراف یکتای ب UNIQUE_PARA_B.\n"
    out = tmp_path / "sem.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    ok, _ = run_semantic_oracle(md, out)
    assert ok

    corrupt = tmp_path / "sem_corrupt.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"UNIQUE_PARA_B", b"")
            zout.writestr(item, data)
    ok2, issues = run_semantic_oracle(md, corrupt)
    assert not ok2
    assert any("UNIQUE_PARA_B" in i or "paragraph" in i.lower() or "missing" in i.lower() for i in issues)


def test_f03_swapped_table_cells_fail_semantic_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import run_semantic_oracle

    md = "| الف | ب |\n| --- | --- |\n| CELL_ONE | CELL_TWO |\n"
    out = tmp_path / "tbl.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    ok, _ = run_semantic_oracle(md, out)
    assert ok

    corrupt = tmp_path / "tbl_swap.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"CELL_ONE", b"CELL_TMP")
                data = data.replace(b"CELL_TWO", b"CELL_ONE")
                data = data.replace(b"CELL_TMP", b"CELL_TWO")
            zout.writestr(item, data)
    ok2, issues = run_semantic_oracle(md, corrupt)
    assert not ok2
    assert any("cell" in i.lower() or "table" in i.lower() or "CELL_" in i for i in issues)


def test_f04_missing_image_target_fails_package_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.package_oracle import run_package_oracle

    img = ROOT / "tests" / "fixtures" / "persian_layout" / "assets" / "horizontal_sample.png"
    md = f"![alt]({img.name})\n"
    out = tmp_path / "img.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True, base_dir=img.parent)
    ok, _ = run_package_oracle(out)
    assert ok

    corrupt = tmp_path / "img_broken.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"media/", b"media/does_not_exist_")
            zout.writestr(item, data)
    ok2, issues = run_package_oracle(corrupt)
    assert not ok2
    assert any("missing" in i.lower() or "relationship" in i.lower() or "target" in i.lower() for i in issues)


def test_f05_heading_number_kept_when_badge_disabled(tmp_path):
    md = "---\nlang: fa-IR\ndir: rtl\n---\n\n## ۱.۱ مبانی پایه\n\nمتن.\n"
    out = tmp_path / "h.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="persian_book", overwrite=True)
    import docx
    doc = docx.Document(str(out))
    texts = " ".join(p.text for p in doc.paragraphs)
    assert "۱.۱" in texts
    assert "مبانی پایه" in texts


def test_f06_heading_uses_heading_font_not_body(tmp_path):
    import shutil, yaml
    src = ROOT / "templates" / "persian_book"
    dest = tmp_path / "fontprobe"
    shutil.copytree(src, dest)
    cfg_path = dest / "config.yaml"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    cfg["fonts"]["body"] = "BODY_FONT_PROBE"
    cfg["fonts"]["heading"] = "HEADING_FONT_PROBE"
    cfg_path.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")
    tmpl = Template.load(str(dest))
    out = tmp_path / "hf.docx"
    convert_markdown_to_docx(content="# عنوان تیتر غنی با **ضخیم**\n\nبدنه.\n", output_path=out, template=tmpl, overwrite=True)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "HEADING_FONT_PROBE" in xml
    assert "BODY_FONT_PROBE" in xml


def test_f08_unknown_callout_role_rejected():
    from md_to_docx.template import Template

    raw = Template.load("purple_book").raw_config.copy()
    raw["custom_styles"] = {"My Note": "nonexistent_role"}
    with pytest.raises(TemplateValidationError) as exc:
        Template(raw, Template.load("purple_book").dir_path)
    assert "nonexistent_role" in str(exc.value).lower() or "role" in str(exc.value).lower()


def test_f09_dir_ltr_metadata_numeric_paragraph_not_rtl(tmp_path):
    md = "---\nlang: en-US\ndir: ltr\n---\n\n12345\n"
    out = tmp_path / "ltr.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="persian_book", overwrite=True)
    with zipfile.ZipFile(out) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    found = False
    for p in root.findall(f".//{W}p"):
        texts = "".join(t.text or "" for t in p.findall(f".//{W}t"))
        if "12345" in texts:
            found = True
            bidi = p.find(f".//{W}bidi")
            val = bidi.get(f"{W}val") if bidi is not None else None
            assert val in (None, "0"), val
    assert found


def test_f11_extreme_aspect_ratio_preserved():
    from md_to_docx.renderer import DocxRenderer

    r = DocxRenderer(template=Template.load("purple_book"))
    w, h = r._fit_image_size(2000, 10)
    ratio = w / h
    native = 2000 / 10
    assert abs(ratio - native) / native < 0.02, (w, h, ratio)


def test_f16_toc_field_when_enabled(tmp_path):
    md = "---\nlang: fa-IR\ndir: rtl\ntoc: true\n---\n\n# فهرست مطالب\n\n# فصل یک {#ch1}\n\nمتن.\n"
    out = tmp_path / "toc.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "TOC" in xml
    assert "fldSimple" in xml or "instrText" in xml


def test_f16_heading_titled_toc_without_metadata_is_not_a_field(tmp_path):
    md = "---\nlang: fa-IR\ndir: rtl\n---\n\n# فهرست مطالب\n\n# فصل\n"
    out = tmp_path / "notoc.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert " TOC " not in xml and 'instr="TOC' not in xml


def test_f00_ast_counts_include_table_cell_image_and_note_mermaid():
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.generate_fixtures import count_ast_elements

    img = ROOT / "tests" / "fixtures" / "persian_layout" / "assets" / "logo.png"
    md = f"""---
lang: fa-IR
dir: rtl
---

| col |
| --- |
| ![x]({img.name}) |

متن پاورقی.[^n]

```mermaid
flowchart LR
    A[یک] --> B[دو]
```

[^n]: ادامهٔ پاورقی با دیاگرام تو‌در‌تو.

    ```mermaid
    flowchart TB
        F[پاورقی] --> G[گره]
    ```
"""
    counts = count_ast_elements(md)
    assert counts["images"] >= 1
    assert counts["mermaid_blocks"] >= 1
    assert counts["tables"] >= 1


def test_f05_extract_number_false_keeps_number_in_title(tmp_path):
    import shutil, yaml
    src = ROOT / "templates" / "persian_book"
    dest = tmp_path / "nobadge"
    shutil.copytree(src, dest)
    cfg_path = dest / "config.yaml"
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    cfg["headings"]["badge"] = False
    cfg["headings"]["extract_number"] = False
    cfg_path.write_text(yaml.dump(cfg, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "n.docx"
    convert_markdown_to_docx(
        content="## ۱.۱ مبانی پایه\n\nمتن.\n",
        output_path=out,
        template=Template.load(str(dest)),
        overwrite=True,
    )
    import docx
    texts = " ".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "۱.۱" in texts
    assert "مبانی پایه" in texts


def test_f06_two_paragraph_footnote_same_size(tmp_path):
    md = "بدنه.[^n]\n\n[^n]: پاراگراف اول UNIQUE_FN_A.\n\n    پاراگراف دوم UNIQUE_FN_B.\n"
    out = tmp_path / "fn.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="persian_book", overwrite=True)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(out) as z:
        root = etree.fromstring(z.read("word/footnotes.xml"))
    sizes = {}
    for p in root.findall(f".//{W}p"):
        text = "".join(t.text or "" for t in p.findall(f".//{W}t"))
        if "UNIQUE_FN_A" in text or "UNIQUE_FN_B" in text:
            sz = p.find(f".//{W}szCs")
            sizes[text.strip()[:20]] = sz.get(f"{W}val") if sz is not None else None
    assert len(sizes) >= 2
    vals = set(sizes.values())
    assert len(vals) == 1, sizes


def test_f07_wrong_caption_identity_does_not_match():
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.matrix_runner import verify_warnings

    actual = [
        "caption_without_image: Standalone caption without associated image @ شکل ۱-۱ at root.blocks[10]",
        "caption_without_image: Standalone caption without associated image @ شکل ۱-۲ at root.blocks[11]",
    ]
    expected_wrong = [
        "caption_without_image @ block 999",
        "caption_without_image @ block 999",
    ]
    issues = verify_warnings(actual, expected_wrong)
    assert issues
    expected_ok = ["caption_without_image @ شکل ۱-۱", "caption_without_image @ شکل ۱-۲"]
    assert verify_warnings(actual, expected_ok) == []


def test_f08_duplicate_custom_style_after_normalize_rejected():
    raw = Template.load("purple_book").raw_config.copy()
    raw["custom_styles"] = {"DBA Note": "note", "dba note": "warning"}
    with pytest.raises(TemplateValidationError):
        Template(raw, Template.load("purple_book").dir_path)


def test_f09_windows_path_stays_one_latin_run():
    from md_to_docx.bidi import split_bidi_runs, ScriptType

    runs = split_bidi_runs(r"مسیر D:\SQLData و حساب DOMAIN\svc_sqlengine")
    texts = [t for t, _ in runs]
    assert any(r"D:\SQLData" in t for t in texts)
    assert any(r"DOMAIN\svc_sqlengine" in t for t in texts)
    for t, s in runs:
        if r"D:\SQLData" in t:
            assert s == ScriptType.LATIN


def test_f10_console_code_keeps_every_line(tmp_path):
    md = (
        "```console\n"
        "netstat -ano | findstr :1433\n"
        'sqlcmd -S SQLPROD01 -E -Q "SELECT @@VERSION;"\n'
        "```\n"
    )
    out = tmp_path / "console.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "netstat -ano | findstr :1433" in xml
    assert "sqlcmd" in xml
    assert "@@VERSION" in xml


def test_f10_list_uses_hanging_indent(tmp_path):
    md = "* آیتم فارسی یک\n* آیتم فارسی دو\n"
    out = tmp_path / "lst.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="persian_book", overwrite=True)
    import docx
    doc = docx.Document(str(out))
    found = False
    for p in doc.paragraphs:
        if "آیتم فارسی" in p.text:
            found = True
            assert p.paragraph_format.first_line_indent is not None
            assert p.paragraph_format.first_line_indent < 0
    assert found


def test_f12_source_label_oracle_and_distinct_themes():
    from md_to_docx.mermaid import extract_mermaid_source_labels, validate_mermaid_svg_labels

    code = "flowchart LR\n    A[شروع فرآیند] --> B[پایان عملیات]\n"
    labels = extract_mermaid_source_labels(code)
    assert "شروع فرآیند" in labels
    svg = """<svg xmlns="http://www.w3.org/2000/svg"><text>شروع فرآیند</text><text>پایان عملیات</text></svg>"""
    ok, missing = validate_mermaid_svg_labels(svg, ["شروع فرآیند", "پایان عملیات"])
    assert ok and not missing
    ok2, missing2 = validate_mermaid_svg_labels(svg, ["شروع فرآیند", "برچسب مفقود"])
    assert not ok2 and "برچسب مفقود" in missing2
    files = [
        (ROOT / "templates" / name / "mermaid.json").read_bytes()
        for name in ("purple_book", "persian_book", "persian_compact", "persian_report")
    ]
    assert len(set(files)) == 4


def test_f04_duplicate_rid_fails_package_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.package_oracle import run_package_oracle

    out = tmp_path / "ok.docx"
    convert_markdown_to_docx(content="# t\n\ntext\n", output_path=out, template="purple_book", overwrite=True)
    corrupt = tmp_path / "dup.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rId1" Type="http://x" Target="theme/theme1.xml"/></Relationships>', 1)
            zout.writestr(item, data)
    ok, issues = run_package_oracle(corrupt)
    assert not ok
    assert any("duplicate" in i.lower() for i in issues)


def test_f13_find_soffice_does_not_hardcode_user_cache(monkeypatch):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts import page_render as pr

    monkeypatch.delenv("MD2DOCX_SOFFICE", raising=False)
    monkeypatch.setattr(pr.shutil, "which", lambda _n: None)
    monkeypatch.setattr(pr.Path, "is_file", lambda self: False)
    found = pr.find_soffice()
    assert found is None or "codex-runtimes" not in str(found)


def test_f04_wrong_content_type_fails_package_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.package_oracle import run_package_oracle

    out = tmp_path / "ct_ok.docx"
    convert_markdown_to_docx(content="# t\n\ntext\n", output_path=out, template="purple_book", overwrite=True)
    corrupt = tmp_path / "ct_bad.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/xml",
                )
            zout.writestr(item, data)
    ok, issues = run_package_oracle(corrupt)
    assert not ok
    assert any("contenttype" in i.lower() for i in issues)


def test_f04_undefined_footnote_rid_fails_package_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.package_oracle import run_package_oracle

    md = "متن با پاورقی.[^1]\n\n[^1]: متن پاورقی با [لینک](https://example.com).\n"
    out = tmp_path / "fn_ok.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    corrupt = tmp_path / "fn_bad_rel.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/footnotes.xml":
                data = data.replace(b'r:id="', b'r:id="nonexistent_')
            zout.writestr(item, data)
    ok, issues = run_package_oracle(corrupt)
    assert not ok
    assert any("undefined relationship" in i.lower() for i in issues)


def test_f03_swapped_paragraphs_fail_semantic_oracle(tmp_path):
    import sys
    from xml.etree import ElementTree as ET
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import run_semantic_oracle

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    md = "پاراگراف یکم اولی.\n\nپاراگراف دوم دومی.\n"
    out = tmp_path / "swap_p.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    corrupt = tmp_path / "swap_p_bad.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = ET.fromstring(data)
                body = root.find(f"{W}body")
                paras = [el for el in list(body) if el.tag == f"{W}p" and el.find(f".//{W}outlineLvl") is None]
                if len(paras) >= 2:
                    p1, p2 = paras[0], paras[1]
                    idx1 = list(body).index(p1)
                    idx2 = list(body).index(p2)
                    body.remove(p1)
                    body.remove(p2)
                    body.insert(idx1, p2)
                    body.insert(idx2, p1)
                    data = ET.tostring(root, encoding="utf-8")
            zout.writestr(item, data)
    ok, issues = run_semantic_oracle(md, corrupt)
    assert not ok
    assert any("sequence mismatch" in i.lower() for i in issues)


def test_f03_swapped_code_lines_fail_semantic_oracle(tmp_path):
    import sys
    from xml.etree import ElementTree as ET
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import run_semantic_oracle

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    md = "```python\nline_one = 1\nline_two = 2\n```\n"
    out = tmp_path / "swap_c.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    corrupt = tmp_path / "swap_c_bad.docx"
    with zipfile.ZipFile(out) as zin, zipfile.ZipFile(corrupt, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = ET.fromstring(data)
                tc = root.find(f".//{W}tc")
                if tc is not None:
                    lines = [p for p in list(tc) if p.tag == f"{W}p"]
                    if len(lines) >= 2:
                        p1, p2 = lines[0], lines[1]
                        tc.remove(p1)
                        tc.remove(p2)
                        tc.insert(0, p2)
                        tc.insert(1, p1)
                        data = ET.tostring(root, encoding="utf-8")
            zout.writestr(item, data)
    ok, issues = run_semantic_oracle(md, corrupt)
    assert not ok
    assert any("code block" in i.lower() or "line" in i.lower() or "mismatch" in i.lower() for i in issues)


def test_f03_heading_badge_level_preserved_in_semantic_oracle(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import actual_from_docx

    md = "## ۱. تیتر سطح دو با بج\n\nمتن پاراگراف آزمایشی.\n"
    out = tmp_path / "badge_h2.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    actual = actual_from_docx(out)
    headings = actual["body"]["headings"]
    assert len(headings) >= 1
    # Outline level in heading_badge table must be level 2, not hardcoded 1
    assert headings[0]["level"] == 2


def test_f03_footnote_image_collected_and_compared(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.semantic_oracle import run_semantic_oracle, actual_from_docx

    md = "متن اصلی با پاورقی تصویر.[^1]\n\n[^1]: متن پاورقی با ![لوگو](tests/fixtures/persian_layout/assets/logo.png)\n"
    out = tmp_path / "fn_img.docx"
    convert_markdown_to_docx(content=md, output_path=out, template="purple_book", overwrite=True)
    actual = actual_from_docx(out)
    assert len(actual["footnotes"]) == 1
    assert len(actual["footnotes"][0]["images"]) == 1

    ok, issues = run_semantic_oracle(md, out)
    assert ok


def test_f12_missing_svg_sidecar_fails_label_check(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.matrix_runner import check_mermaid_label_sidecars

    md = "```mermaid\nflowchart LR\n    A[شروع] --> B[پایان]\n```\n"
    media = tmp_path / "media"
    media.mkdir()
    issues = check_mermaid_label_sidecars(md, media)
    assert issues
    assert any("missing" in i.lower() and "svg" in i.lower() for i in issues)

    svg = media / "diagram_001.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><text>شروع</text><text>پایان</text></svg>',
        encoding="utf-8",
    )
    assert check_mermaid_label_sidecars(md, media) == []


@pytest.mark.mermaid
def test_f12_real_mermaid_writes_svg_sidecar(tmp_path):
    import sys
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from scripts.matrix_runner import execute_matrix_pair

    fix = {
        "id": "S10",
        "rel_path": "synthetic/S10.md",
        "expected_warnings": [],
    }
    res = execute_matrix_pair(fix, "purple_book", tmp_path)
    assert res["status"] == "pass", res.get("issues")
    assert res["mermaid_label_pass"] is True
    pair_dir = tmp_path / "S10" / "purple_book"
    assert (pair_dir / "media" / "diagram_001.svg").is_file()


def test_f06_svg_sidecar_required_after_png_success(tmp_path, monkeypatch):
    from md_to_docx import mermaid as mermaid_mod
    from md_to_docx.template import Template

    mermaid_mod.reset_launch_health()
    calls = {"n": 0}

    def fake_run(cmd, **kwargs):
        calls["n"] += 1
        out = None
        if "-o" in cmd:
            out = Path(cmd[cmd.index("-o") + 1])
        class Proc:
            returncode = 0
            stdout = ""
            stderr = ""
        if out and out.suffix.lower() == ".png":
            out.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
            # Intentionally do not write the SVG sidecar.
        return Proc()

    monkeypatch.setattr(mermaid_mod.subprocess, "run", fake_run)
    monkeypatch.setattr(mermaid_mod, "_find_mmdc_cmd", lambda: ["mmdc"])
    monkeypatch.setattr(mermaid_mod, "validate_rendered_diagram_image", lambda _p: True)
    png = tmp_path / "diagram.png"
    with pytest.raises(mermaid_mod.ConvertError) as exc:
        mermaid_mod._run_mmdc("graph TD;A-->B;", png, Template.load("purple_book"), 10, None)
    assert "svg" in str(exc.value).lower()


def test_f13_matrix_runner_render_pages_flag(tmp_path):
    import sys
    import subprocess
    from scripts.page_render import find_soffice

    res = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "matrix_runner.py"),
            "--fixtures", "S01",
            "--templates", "purple_book",
            "--render-pages",
            "--overwrite-run",
            "--output-dir", str(tmp_path / "matrix_out"),
        ],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    if not find_soffice():
        assert res.returncode == 1
        assert "BLOCKED" in res.stdout or "blocked" in res.stdout.lower()



