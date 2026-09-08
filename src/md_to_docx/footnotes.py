"""Create real Word footnotes (word/footnotes.xml) rather than inline superscript text."""

from __future__ import annotations

from typing import Any, Optional
from docx.document import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsmap, qn
from docx.oxml.shape import CT_Inline
from docx.text.paragraph import Paragraph

W_NS = nsmap["w"]
FOOTNOTES_URI = "/word/footnotes.xml"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"


def _nsmap() -> dict:
    return {"w": W_NS}


def _wire_part_image_methods(part: Part, package: Any) -> None:
    """Wire image and inline pic helpers onto the footnotes Part so runs can embed images."""
    if hasattr(part, "new_pic_inline"):
        return

    def get_or_add_image(image_descriptor):
        image_part = package.get_or_add_image_part(image_descriptor)
        rId = part.relate_to(image_part, RT.IMAGE)
        return rId, image_part.image

    def new_pic_inline(image_descriptor, width=None, height=None):
        rId, image = get_or_add_image(image_descriptor)
        cx, cy = image.scaled_dimensions(width, height)
        shape_id = getattr(part, "_next_shape_id", 1)
        part._next_shape_id = shape_id + 1
        return CT_Inline.new_pic_inline(shape_id, rId, image.filename, cx, cy)

    part.get_or_add_image = get_or_add_image
    part.new_pic_inline = new_pic_inline


def _ensure_footnotes_part(doc: Document) -> Part:
    rels = doc.part.rels
    for rel in rels.values():
        if rel.reltype == RT.FOOTNOTES:
            part = rel.target_part
            _wire_part_image_methods(part, doc.part.package)
            return part

    xml = (
        f'<w:footnotes xmlns:w="{W_NS}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<w:footnote w:type="separator" w:id="-1">'
        "<w:p><w:r><w:separator/></w:r></w:p>"
        "</w:footnote>"
        '<w:footnote w:type="continuationSeparator" w:id="0">'
        "<w:p><w:r><w:continuationSeparator/></w:r></w:p>"
        "</w:footnote>"
        "</w:footnotes>"
    )
    part = Part(
        PackURI(FOOTNOTES_URI),
        FOOTNOTES_CT,
        xml.encode("utf-8"),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)
    try:
        ct = doc.part.package.content_types
        ct._overrides[FOOTNOTES_URI] = FOOTNOTES_CT
    except Exception:
        pass
    _wire_part_image_methods(part, doc.part.package)
    return part


def next_footnote_id(part: Part) -> int:
    root = part._element if hasattr(part, "_element") else parse_xml(part.blob)
    ids = [int(el.get(qn("w:id"))) for el in root.findall(qn("w:footnote")) if el.get(qn("w:id"))]
    return max(ids) + 1 if ids else 1


def add_footnote_reference(paragraph: Paragraph, footnote_id: int) -> None:
    r = paragraph.add_run()
    rPr = r._r.get_or_add_rPr()
    vert = OxmlElement("w:vertAlign")
    vert.set(qn("w:val"), "superscript")
    rPr.append(vert)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(footnote_id))
    r._r.append(ref)


class FootnoteContainer:
    """Container for footnote body blocks, managing paragraph creation and part relationships."""

    def __init__(self, doc: Document, part: Part, root_el: Any, fn_el: Any, first_p_el: Any):
        self._doc = doc
        self.part = part
        self._root = root_el
        self._fn_element = fn_el
        self._first_p = Paragraph(first_p_el, self)
        self.paragraphs = [self._first_p]

    @property
    def first_paragraph(self) -> Paragraph:
        return self._first_p

    def add_paragraph(self, text: str = "") -> Paragraph:
        p = OxmlElement("w:p")
        self._fn_element.append(p)
        p_obj = Paragraph(p, self)
        if text:
            p_obj.text = text
        self.paragraphs.append(p_obj)
        return p_obj

    def add_table(self, rows: int, cols: int, style=None):
        """Append a table to the footnote body (tables/code/callouts in notes).

        Mirrors ``_Cell.add_table``: the table is appended to the ``w:footnote``
        element and wrapped so ``tbl.cell()`` and ``part`` resolve to the
        footnotes part. Column widths are set later by the table renderer.
        """
        from docx.oxml.table import CT_Tbl
        from docx.shared import Inches
        from docx.table import Table

        tbl_el = CT_Tbl.new_tbl(rows, cols, Inches(6))
        self._fn_element.append(tbl_el)
        tbl = Table(tbl_el, self)
        if style is not None:
            tbl.style = style
        return tbl

    @classmethod
    def create(cls, doc: Document, part: Part, footnote_id: int) -> tuple[FootnoteContainer, Any]:
        _wire_part_image_methods(part, doc.part.package)
        root = parse_xml(part.blob)
        fn = OxmlElement("w:footnote")
        fn.set(qn("w:id"), str(footnote_id))
        p = OxmlElement("w:p")
        r_mark = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rPr.append(vert)
        r_mark.append(rPr)
        r_mark.append(OxmlElement("w:footnoteRef"))
        p.append(r_mark)

        r_space = OxmlElement("w:r")
        t_space = OxmlElement("w:t")
        t_space.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t_space.text = " "
        r_space.append(t_space)
        p.append(r_space)

        fn.append(p)
        root.append(fn)
        container = cls(doc, part, root, fn, p)
        return container, root


def add_footnote_body(doc: Document, footnote_id: int):
    """Append an empty footnote paragraph and return (paragraph, root, part) for flushing."""
    part = _ensure_footnotes_part(doc)
    container, root = FootnoteContainer.create(doc, part, footnote_id)
    part._blob = ET_to_bytes(root)
    return container.first_paragraph, root, part


def ET_to_bytes(root) -> bytes:
    from lxml import etree
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def list_footnote_ids(doc: Document) -> list[int]:
    try:
        part = None
        for rel in doc.part.rels.values():
            if rel.reltype == RT.FOOTNOTES:
                part = rel.target_part
                break
        if part is None:
            return []
        root = parse_xml(part.blob)
        out = []
        for el in root.findall(qn("w:footnote")):
            fid = el.get(qn("w:id"))
            if fid and int(fid) > 0:
                out.append(int(fid))
        return out
    except Exception:
        return []
