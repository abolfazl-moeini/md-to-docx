#!/usr/bin/env python3
"""
Sets up and synchronizes the four templates:
- T01: purple_book
- T02: persian_book
- T03: persian_compact
- T04: persian_report
both in templates/ and src/md_to_docx/templates/.
"""

import shutil
from pathlib import Path
from typing import Optional
import docx
import yaml

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES_ROOT = ROOT / "templates"
PKG_TEMPLATES_ROOT = ROOT / "src" / "md_to_docx" / "templates"

COMMON_CUSTOM_STYLES = {
    "chapter overview": "overview",
    "dba note": "note",
    "important note": "important",
    "warning": "warning",
    "lab note": "lab",
    "screenshot recommendation": "editorial",
}


def create_shell_docx(
    dest: Path,
    page_size: str = "A4",
    has_page_number: bool = True,
    logo_path: Optional[Path] = None,
):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()
    section = doc.sections[0]

    dim_map = {
        "A4": (docx.shared.Cm(21.0), docx.shared.Cm(29.7)),
        "A5": (docx.shared.Cm(14.8), docx.shared.Cm(21.0)),
        "Letter": (docx.shared.Cm(21.59), docx.shared.Cm(27.94)),
    }
    w, h = dim_map.get(page_size, dim_map["A4"])
    section.page_width = w
    section.page_height = h

    if logo_path and logo_path.exists():
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hrun = hp.add_run()
        hrun.add_picture(str(logo_path), width=docx.shared.Inches(0.5))

    if has_page_number:
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = "صفحه "
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), " PAGE ")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        fp._p.append(fld)

    doc.save(str(dest))


def main():
    font_dir = TEMPLATES_ROOT / "purple_book" / "fonts"

    # 1. Setup T01: purple_book
    for base in [TEMPLATES_ROOT, PKG_TEMPLATES_ROOT]:
        pb_dir = base / "purple_book"
        pb_dir.mkdir(parents=True, exist_ok=True)
        config_path = pb_dir / "config.yaml"
        with open(config_path, "r", encoding="utf-8") as f:
            cfg = yaml.safe_load(f)

        cfg["custom_styles"] = dict(COMMON_CUSTOM_STYLES)
        cfg["font_files"] = {
            "Vazirmatn": "fonts/Vazirmatn-Regular.ttf",
            "Vazirmatn-Bold": "fonts/Vazirmatn-Bold.ttf",
        }
        cfg["page"]["paragraph_align"] = "both"
        cfg["page"]["space_after_pt"] = 6.0
        cfg["caption"] = {"size_pt": 9.5}

        with open(config_path, "w", encoding="utf-8") as f:
            yaml.dump(cfg, f, allow_unicode=True, sort_keys=False)

    # 2. Setup T02: persian_book (A4)
    for base in [TEMPLATES_ROOT, PKG_TEMPLATES_ROOT]:
        tb_dir = base / "persian_book"
        tb_dir.mkdir(parents=True, exist_ok=True)
        shutil.copytree(font_dir, tb_dir / "fonts", dirs_exist_ok=True)
        for f in ["mermaid.json", "mermaid.css", "puppeteer.json"]:
            shutil.copy2(TEMPLATES_ROOT / "purple_book" / f, tb_dir / f)

        create_shell_docx(tb_dir / "shell.docx", page_size="A4", has_page_number=True)

        cfg_t02 = {
            "schema_version": 1,
            "name": "persian_book",
            "direction": "rtl",
            "language_bidi": "fa-IR",
            "language_latin": "en-US",
            "fonts": {
                "body": "Vazirmatn",
                "heading": "Vazirmatn",
                "latin": "Vazirmatn",
                "code": "Courier New",
            },
            "font_files": {
                "Vazirmatn": "fonts/Vazirmatn-Regular.ttf",
                "Vazirmatn-Bold": "fonts/Vazirmatn-Bold.ttf",
            },
            "shell": "shell.docx",
            "page": {
                "size": "A4",
                "margin_cm": {"top": 2.2, "bottom": 2.2, "left": 2.2, "right": 2.2},
                "font_size_pt": 12.0,
                "line_spacing": 1.35,
                "paragraph_align": "start",
                "space_after_pt": 6.0,
            },
            "caption": {
                "size_pt": 10.0,
            },
            "custom_styles": dict(COMMON_CUSTOM_STYLES),
            "colors": {
                "primary": "17324D",
                "primary_dark": "0F2133",
                "on_primary": "FFFFFF",
                "quote_bg": "E8EEF3",
                "warning_bg": "FBF7F4",
                "warning_title": "8B6914",
                "body": "222222",
                "caption": "555555",
            },
            "headings": {
                "extract_number": True,
                "badge": False,
                "h1": {"size_pt": 22, "page_break_before": True},
                "h2": {"size_pt": 18},
                "h3": {"size_pt": 15},
                "h4": {"size_pt": 13},
                "h5": {"size_pt": 12},
                "h6": {"size_pt": 12},
            },
            "callouts": {
                "note": {
                    "classes": ["note"],
                    "default_title": "نکته",
                    "icon": "◆",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "F0F4F8",
                },
                "warning": {
                    "classes": ["warning"],
                    "default_title": "هشدار",
                    "icon": "",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "warning_bg",
                },
                "overview": {
                    "classes": ["overview"],
                    "default_title": "",
                    "icon": "",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "E8EEF3",
                },
                "important": {
                    "classes": ["important"],
                    "default_title": "توجه مهم",
                    "icon": "!",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "FDF8F0",
                },
                "lab": {
                    "classes": ["lab"],
                    "default_title": "یادداشت تمرین",
                    "icon": "◆",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "E8EEF3",
                },
                "editorial": {
                    "classes": ["editorial"],
                    "default_title": "پیشنهاد تصویر",
                    "icon": "◆",
                    "header_bg": "E8EEF3",
                    "header_fg": "caption",
                    "body_bg": "F8F9FA",
                },
            },
            "quotes": {
                "border_side": "physical_right",
                "border_pt": 12,
                "border_color": "primary",
                "bg": "quote_bg",
            },
            "tables": {
                "header_bg": "primary",
                "header_fg": "on_primary",
                "bidi_visual": True,
            },
            "code_block": {
                "theme": "friendly",
                "bg": "F4F6F8",
                "border_color": "CFD8DC",
                "border_sz": 4,
                "font_size_pt": 9.5,
                "line_spacing": 1.15,
            },
            "mermaid": {
                "format": "png",
                "scale": 3,
                "max_width_in": 6.3,
                "theme_file": "mermaid.json",
                "css_file": "mermaid.css",
                "puppeteer_config": "puppeteer.json",
            },
        }
        with open(tb_dir / "config.yaml", "w", encoding="utf-8") as f:
            yaml.dump(cfg_t02, f, allow_unicode=True, sort_keys=False)

    # 3. Setup T03: persian_compact (A5)
    for base in [TEMPLATES_ROOT, PKG_TEMPLATES_ROOT]:
        tc_dir = base / "persian_compact"
        tc_dir.mkdir(parents=True, exist_ok=True)
        shutil.copytree(font_dir, tc_dir / "fonts", dirs_exist_ok=True)
        for f in ["mermaid.json", "mermaid.css", "puppeteer.json"]:
            shutil.copy2(TEMPLATES_ROOT / "purple_book" / f, tc_dir / f)

        create_shell_docx(tc_dir / "shell.docx", page_size="A5", has_page_number=True)

        cfg_t03 = {
            "schema_version": 1,
            "name": "persian_compact",
            "direction": "rtl",
            "language_bidi": "fa-IR",
            "language_latin": "en-US",
            "fonts": {
                "body": "Vazirmatn",
                "heading": "Vazirmatn",
                "latin": "Vazirmatn",
                "code": "Courier New",
            },
            "font_files": {
                "Vazirmatn": "fonts/Vazirmatn-Regular.ttf",
                "Vazirmatn-Bold": "fonts/Vazirmatn-Bold.ttf",
            },
            "shell": "shell.docx",
            "page": {
                "size": "A5",
                "margin_cm": {"top": 1.6, "bottom": 1.6, "left": 1.6, "right": 1.6},
                "font_size_pt": 11.5,
                "line_spacing": 1.25,
                "paragraph_align": "start",
                "space_after_pt": 5.0,
            },
            "caption": {
                "size_pt": 9.0,
            },
            "custom_styles": dict(COMMON_CUSTOM_STYLES),
            "colors": {
                "primary": "2E3440",
                "primary_dark": "242933",
                "on_primary": "FFFFFF",
                "quote_bg": "ECEFF4",
                "warning_bg": "FBF7F4",
                "warning_title": "8B6914",
                "body": "2E3440",
                "caption": "4C566A",
            },
            "headings": {
                "extract_number": True,
                "badge": False,
                "h1": {"size_pt": 19, "page_break_before": True},
                "h2": {"size_pt": 16},
                "h3": {"size_pt": 13.5},
                "h4": {"size_pt": 12},
                "h5": {"size_pt": 11.5},
                "h6": {"size_pt": 11.5},
            },
            "callouts": {
                "note": {
                    "classes": ["note"],
                    "default_title": "نکته",
                    "icon": "◆",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "ECEFF4",
                },
                "warning": {
                    "classes": ["warning"],
                    "default_title": "هشدار",
                    "icon": "",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "warning_bg",
                },
                "overview": {
                    "classes": ["overview"],
                    "default_title": "",
                    "icon": "",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "ECEFF4",
                },
                "important": {
                    "classes": ["important"],
                    "default_title": "توجه مهم",
                    "icon": "!",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "F9F9F9",
                },
                "lab": {
                    "classes": ["lab"],
                    "default_title": "یادداشت تمرین",
                    "icon": "◆",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "ECEFF4",
                },
                "editorial": {
                    "classes": ["editorial"],
                    "default_title": "پیشنهاد تصویر",
                    "icon": "◆",
                    "header_bg": "ECEFF4",
                    "header_fg": "caption",
                    "body_bg": "FAFAFA",
                },
            },
            "quotes": {
                "border_side": "physical_right",
                "border_pt": 8,
                "border_color": "primary",
                "bg": "quote_bg",
            },
            "tables": {
                "header_bg": "primary",
                "header_fg": "on_primary",
                "bidi_visual": True,
            },
            "code_block": {
                "theme": "friendly",
                "bg": "F8F9FA",
                "border_color": "D8DEE9",
                "border_sz": 4,
                "font_size_pt": 9.0,
                "line_spacing": 1.15,
            },
            "mermaid": {
                "format": "png",
                "scale": 3,
                "max_width_in": 4.5,
                "theme_file": "mermaid.json",
                "css_file": "mermaid.css",
                "puppeteer_config": "puppeteer.json",
            },
        }
        with open(tc_dir / "config.yaml", "w", encoding="utf-8") as f:
            yaml.dump(cfg_t03, f, allow_unicode=True, sort_keys=False)

    # 4. Setup T04: persian_report (Letter)
    for base in [TEMPLATES_ROOT, PKG_TEMPLATES_ROOT]:
        tr_dir = base / "persian_report"
        tr_dir.mkdir(parents=True, exist_ok=True)
        shutil.copytree(font_dir, tr_dir / "fonts", dirs_exist_ok=True)
        for f in ["mermaid.json", "mermaid.css", "puppeteer.json"]:
            shutil.copy2(TEMPLATES_ROOT / "purple_book" / f, tr_dir / f)

        logo_file = tr_dir / "assets" / "logo.png"
        create_shell_docx(tr_dir / "shell.docx", page_size="Letter", has_page_number=True, logo_path=logo_file)

        cfg_t04 = {
            "schema_version": 1,
            "name": "persian_report",
            "direction": "rtl",
            "language_bidi": "fa-IR",
            "language_latin": "en-US",
            "fonts": {
                "body": "Vazirmatn",
                "heading": "Vazirmatn",
                "latin": "Vazirmatn",
                "code": "Courier New",
            },
            "font_files": {
                "Vazirmatn": "fonts/Vazirmatn-Regular.ttf",
                "Vazirmatn-Bold": "fonts/Vazirmatn-Bold.ttf",
            },
            "shell": "shell.docx",
            "page": {
                "size": "Letter",
                "margin_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0},
                "font_size_pt": 11.5,
                "line_spacing": 1.25,
                "paragraph_align": "start",
                "space_after_pt": 6.0,
            },
            "caption": {
                "size_pt": 10.0,
            },
            "custom_styles": dict(COMMON_CUSTOM_STYLES),
            "colors": {
                "primary": "17324D",
                "primary_dark": "006D77",
                "on_primary": "FFFFFF",
                "quote_bg": "E2F1F1",
                "warning_bg": "FDF8F0",
                "warning_title": "8B6914",
                "body": "24292E",
                "caption": "586069",
            },
            "headings": {
                "extract_number": True,
                "badge": False,
                "h1": {"size_pt": 20},
                "h2": {"size_pt": 17},
                "h3": {"size_pt": 14},
                "h4": {"size_pt": 12.5},
                "h5": {"size_pt": 12},
                "h6": {"size_pt": 11.5},
            },
            "callouts": {
                "note": {
                    "classes": ["note"],
                    "default_title": "نکته",
                    "icon": "◆",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "E2F1F1",
                },
                "warning": {
                    "classes": ["warning"],
                    "default_title": "هشدار",
                    "icon": "",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "warning_bg",
                },
                "overview": {
                    "classes": ["overview"],
                    "default_title": "",
                    "icon": "",
                    "header_bg": "primary_dark",
                    "header_fg": "on_primary",
                    "body_bg": "E2F1F1",
                },
                "important": {
                    "classes": ["important"],
                    "default_title": "توجه مهم",
                    "icon": "!",
                    "header_bg": "warning_bg",
                    "header_fg": "warning_title",
                    "body_bg": "FDF8F0",
                },
                "lab": {
                    "classes": ["lab"],
                    "default_title": "یادداشت تمرین",
                    "icon": "◆",
                    "header_bg": "primary",
                    "header_fg": "on_primary",
                    "body_bg": "E2F1F1",
                },
                "editorial": {
                    "classes": ["editorial"],
                    "default_title": "پیشنهاد تصویر",
                    "icon": "◆",
                    "header_bg": "E2F1F1",
                    "header_fg": "caption",
                    "body_bg": "F8F9FA",
                },
            },
            "quotes": {
                "border_side": "physical_right",
                "border_pt": 12,
                "border_color": "primary_dark",
                "bg": "quote_bg",
            },
            "tables": {
                "header_bg": "primary",
                "header_fg": "on_primary",
                "bidi_visual": True,
            },
            "code_block": {
                "theme": "friendly",
                "bg": "F6F8FA",
                "border_color": "D0D7DE",
                "border_sz": 4,
                "font_size_pt": 9.5,
                "line_spacing": 1.15,
            },
            "mermaid": {
                "format": "png",
                "scale": 3,
                "max_width_in": 6.3,
                "theme_file": "mermaid.json",
                "css_file": "mermaid.css",
                "puppeteer_config": "puppeteer.json",
            },
        }
        with open(tr_dir / "config.yaml", "w", encoding="utf-8") as f:
            yaml.dump(cfg_t04, f, allow_unicode=True, sort_keys=False)

    print("All 4 templates initialized and synchronized.")


if __name__ == "__main__":
    main()
